from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from corporate_style import LOGO_SHA256, NAVY, STYLE_VERSION, _build_self_test_workbook, apply_corporate_a4


def main() -> int:
    styled = apply_corporate_a4(_build_self_test_workbook())
    wb = load_workbook(BytesIO(styled), data_only=False)
    ws = wb.active

    # openpyxl serializes paperSize as an integer even though the class constant is a string.
    assert int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A4)
    assert int(ws.page_setup.fitToWidth) == 1
    assert int(ws.page_setup.fitToHeight) == 1
    assert "$A$1:$G$17" in str(ws.print_area)
    assert len(ws._images) == 1
    assert ws["G14"].value == "=SUM(G12:G12)"
    assert ws["A11"].fill.fgColor.rgb.endswith(NAVY)

    print(
        {
            "status": "CORPORATE_A4_VALIDATION_OK",
            "style_version": STYLE_VERSION,
            "logo_sha256": LOGO_SHA256,
            "paper_size": ws.page_setup.paperSize,
            "fit_to_width": ws.page_setup.fitToWidth,
            "fit_to_height": ws.page_setup.fitToHeight,
            "print_area": str(ws.print_area),
            "images": len(ws._images),
        }
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
