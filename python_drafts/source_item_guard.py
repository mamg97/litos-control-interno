from __future__ import annotations

from io import BytesIO
import re

import openpyxl
from openpyxl.cell.cell import MergedCell

# Only source-explicit absences qualify. Missing source fields are not absences.
_SOURCE_ABSENCES = {
    "CORONACIÓN": ("CORONACION",),
    "COLUMNA / PILASTRA": ("COLUMNAS", "PILASTRAS"),
    "CRUZ": ("CRUZ",),
    "FLORERO": ("FLORERO",),
    "CORNISA": ("CORNISA",),
    "TABICA / TACOS": ("TABICA Y TACOS",),
}


def _explicitly_absent(specifications: str, field: str) -> bool:
    pattern = r"(?:^|[.;\n])\s*" + re.escape(field) + r"\s*:\s*([^.;\n]*)"
    matches = re.findall(pattern, specifications, flags=re.IGNORECASE)
    return len(matches) == 1 and matches[0].strip() == "-"


def _clean_absence_warning(value: str, absent: set[str]) -> str:
    """Remove only source-confirmed absent items from draft warning lists."""
    if not value.startswith("REVISAR DATOS"):
        return value
    parts = value.split(" · ")
    kept_parts = [parts[0]]
    for part in parts[1:]:
        if ":" not in part:
            kept_parts.append(part)
            continue
        label, items = part.split(":", 1)
        if label.strip().casefold() not in {"sin precio automático", "faltan medidas para"}:
            kept_parts.append(part)
            continue
        kept = [
            item.strip() for item in items.split(",")
            if item.strip() and item.strip().upper() not in absent
        ]
        if kept:
            kept_parts.append(f"{label}: {', '.join(kept)}")
    return " · ".join(kept_parts) if len(kept_parts) > 1 else "BORRADOR · REVISAR ANTES DE USAR"


def has_stale_absence_warnings(xlsx_bytes: bytes, absent_labels: list[str]) -> bool:
    """Detect a stale warning even when all unrequested charge rows were cleared."""
    if not absent_labels:
        return False
    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True)
    try:
        absent = {label.upper() for label in absent_labels}
        for row in workbook.active.iter_rows(min_row=1, max_row=12, max_col=7, values_only=True):
            for value in row:
                if isinstance(value, str) and _clean_absence_warning(value, absent) != value:
                    return True
        return False
    finally:
        workbook.close()


def apply_explicit_absences(xlsx_bytes: bytes, source_row: dict) -> tuple[bytes, list[str]]:
    """Clear only template lines explicitly marked absent in the private source."""
    specifications = str(source_row.get("Especificaciones", "") or "")
    absent = {
        label
        for label, keys in _SOURCE_ABSENCES.items()
        if all(_explicitly_absent(specifications, key) for key in keys)
    }
    if not absent:
        return xlsx_bytes, []
    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    removed = []
    try:
        sheet = workbook.active
        for row_idx in range(1, sheet.max_row + 1):
            label = str(sheet.cell(row_idx, 1).value or "").strip().upper()
            if label not in absent:
                continue
            if any(sheet.cell(row_idx, col).value is not None for col in range(1, 8)):
                for col in range(1, 8):
                    cell = sheet.cell(row_idx, col)
                    # Merged placeholders are read-only; clear only real anchor cells.
                    # Never unmerge or shift other source rows to satisfy a draft.
                    if not isinstance(cell, MergedCell):
                        cell.value = None
                removed.append(label)
        # Template warnings are generated before source-absence cleanup.
        # They must not demand measurements/prices for expressly absent items.
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), max_col=7):
            for cell in row:
                if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
                    continue
                cleaned = _clean_absence_warning(cell.value, absent)
                if cleaned != cell.value:
                    cell.value = cleaned
        if not removed:
            return xlsx_bytes, []
        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), sorted(set(removed))
    finally:
        workbook.close()


def has_unrequested_rows(xlsx_bytes: bytes, removed_labels: list[str]) -> bool:
    """Identify an existing draft that still contains explicitly absent items."""
    if not removed_labels:
        return False
    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True)
    try:
        labels = set(removed_labels)
        for row in workbook.active.iter_rows(min_col=1, max_col=7, values_only=True):
            if str(row[0] or "").strip().upper() in labels and any(value is not None for value in row):
                return True
        return False
    finally:
        workbook.close()
