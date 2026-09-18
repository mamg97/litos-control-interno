from __future__ import annotations

from copy import copy
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

CUSTOMER_PROFILE_VERSION = "private-master-client-v1"
REQUIRED_CUSTOMER_KEYS = ("name", "address", "city", "province")

NAVY = "16227A"
DARK = "1F2937"
MUTED = "667085"
LINE = "D6DCE3"
WHITE = "FFFFFF"
RED = "D9272E"

THIN = Side(style="thin", color=LINE)
NAVY_SIDE = Side(style="medium", color=NAVY)
RED_SIDE = Side(style="medium", color=RED)
SOFT_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _text(value) -> str:
    return str(value if value is not None else "").strip()


def _find_title_row(ws) -> int:
    for row in range(1, ws.max_row + 1):
        if _text(ws.cell(row, 1).value).upper() == "ALBARÁN":
            return row
    raise RuntimeError("Client header contract changed: ALBARÁN row not found")


def _transaction_state(ws, start_row: int) -> tuple:
    out = []
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            out.append((cell.coordinate, cell.data_type, cell.value))
    return tuple(out)


def _header_values(ws) -> dict:
    return {
        "company": _text(ws["A1"].value),
        "nif": _text(ws["A2"].value),
        "address": _text(ws["A3"].value),
        "phone": _text(ws["F3"].value),
        "city": _text(ws["A4"].value),
        "trade": _text(ws["A5"].value),
    }


def _unmerge_header(ws) -> None:
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= 1 and merged.max_row <= 5:
            ws.unmerge_cells(str(merged))


def _merge_value(ws, cell_range: str, value: str, *, font: Font, alignment: Alignment,
                 fill: str = WHITE, border: Border | None = None) -> None:
    ws.merge_cells(cell_range)
    top_left = ws[cell_range.split(":", 1)[0]]
    top_left.value = value
    top_left.font = copy(font)
    top_left.alignment = copy(alignment)
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill)
            if border:
                cell.border = copy(border)


def apply_client_header(xlsx_bytes: bytes, customer: dict | None = None) -> bytes:
    """Add the customer identity block while preserving the draft body exactly.

    Rows from ALBARÁN downwards are treated as transactional content and are
    byte-for-byte equivalent at cell value/formula level before and after this
    operation. Only the issuer/customer header area (rows 1:5) is rebuilt.
    """
    if customer is None:
        raise RuntimeError("Customer profile must be supplied from the private master")
    profile = {key: _text(customer.get(key)) for key in REQUIRED_CUSTOMER_KEYS}
    if any(not profile[key] for key in REQUIRED_CUSTOMER_KEYS):
        raise RuntimeError("Customer header requires name, address, city and province")

    wb = load_workbook(BytesIO(xlsx_bytes), data_only=False)
    ws = wb.worksheets[0]
    title_row = _find_title_row(ws)
    if title_row < 6:
        raise RuntimeError("Client header contract changed: expected five-row issuer header")

    before = _transaction_state(ws, title_row)
    issuer = _header_values(ws)
    if not issuer["company"] or not issuer["trade"]:
        raise RuntimeError("Client header contract changed: issuer header is incomplete")

    _unmerge_header(ws)
    for row in range(1, 6):
        for col in range(1, 8):
            cell = ws.cell(row, col)
            cell.value = None
            cell.fill = PatternFill("solid", fgColor=WHITE)
            cell.border = Border()
            cell.alignment = Alignment(vertical="center")

    # Keep column A visually reserved for the immutable corporate logo already
    # embedded by corporate_style.py. Issuer identity occupies B:D.
    _merge_value(
        ws, "B1:D1", issuer["company"],
        font=Font(name="Arial", size=11.5, bold=True, color=NAVY),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    _merge_value(
        ws, "B2:D2", issuer["nif"],
        font=Font(name="Arial", size=8.2, color=MUTED),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    _merge_value(
        ws, "B3:C3", issuer["address"],
        font=Font(name="Arial", size=8.2, color=DARK),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    ws["D3"] = issuer["phone"]
    ws["D3"].font = Font(name="Arial", size=7.5, color=DARK)
    ws["D3"].alignment = Alignment(horizontal="right", vertical="center", shrink_to_fit=True)
    _merge_value(
        ws, "B4:D4", issuer["city"],
        font=Font(name="Arial", size=8.5, bold=True, color=DARK),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    _merge_value(
        ws, "B5:D5", issuer["trade"],
        font=Font(name="Arial", size=8.8, bold=True, color=NAVY),
        alignment=Alignment(horizontal="left", vertical="center"),
    )

    # Customer block: fixed for the current operating model, but the function is
    # parameterized so future orders can supply a different customer cleanly.
    _merge_value(
        ws, "E1:G1", "CLIENTE",
        font=Font(name="Arial", size=8.5, bold=True, color=WHITE),
        alignment=Alignment(horizontal="left", vertical="center"),
        fill=NAVY,
        border=Border(left=NAVY_SIDE, right=NAVY_SIDE, top=NAVY_SIDE, bottom=THIN),
    )
    customer_rows = (
        (2, profile["name"], True),
        (3, profile["address"], False),
        (4, profile["city"], False),
        (5, profile["province"], False),
    )
    for row, value, bold in customer_rows:
        bottom = RED_SIDE if row == 5 else THIN
        _merge_value(
            ws, f"E{row}:G{row}", value,
            font=Font(name="Arial", size=8.4 if row != 2 else 8.6, bold=bold, color=DARK),
            alignment=Alignment(horizontal="left", vertical="center", shrink_to_fit=True),
            border=Border(left=NAVY_SIDE, right=NAVY_SIDE, top=THIN, bottom=bottom),
        )

    # Preserve the compact A4 header height used by the approved corporate style.
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 25)
    for row, minimum in ((2, 16), (3, 18), (4, 18), (5, 18)):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, minimum)

    after = _transaction_state(ws, title_row)
    if after != before:
        raise RuntimeError("Customer header guard: transactional cell values/formulas changed")

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _build_self_test_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:G1")
    ws["A1"] = "EMPRESA DE PRUEBA"
    ws.merge_cells("A2:G2")
    ws["A2"] = "N.I.F. TEST000000"
    ws.merge_cells("A3:E3")
    ws["A3"] = "CALLE PRUEBA 1"
    ws.merge_cells("F3:G3")
    ws["F3"] = "Tlf. 000 000 000"
    ws.merge_cells("A4:G4")
    ws["A4"] = "CIUDAD DE PRUEBA"
    ws.merge_cells("A5:G5")
    ws["A5"] = "MÁRMOLES, PIEDRAS Y GRANITOS"
    ws.merge_cells("A6:G6")
    ws["A6"] = "ALBARÁN"
    ws["A7"] = "BORRADOR"
    ws["A8"] = "PEDIDO Nº"
    ws["B8"] = 9999
    ws["G9"] = "=1+1"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def self_test() -> dict:
    customer = {
        "name": "CLIENTE DE PRUEBA",
        "address": "DIRECCION DE PRUEBA",
        "city": "CIUDAD DE PRUEBA",
        "province": "PROVINCIA DE PRUEBA",
    }
    result = apply_client_header(_build_self_test_workbook(), customer)
    wb = load_workbook(BytesIO(result), data_only=False)
    ws = wb.active
    assert ws["E1"].value == "CLIENTE"
    assert ws["E2"].value == customer["name"]
    assert ws["E3"].value == customer["address"]
    assert ws["E4"].value == customer["city"]
    assert ws["E5"].value == customer["province"]
    assert ws["A6"].value == "ALBARÁN"
    assert ws["G9"].value == "=1+1"
    return {
        "status": "CLIENT_HEADER_VALIDATION_OK",
        "customer_profile_version": CUSTOMER_PROFILE_VERSION,
        "customer": customer,
    }


if __name__ == "__main__":
    print(self_test())
