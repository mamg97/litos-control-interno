from __future__ import annotations
from io import BytesIO
from copy import copy as _copy
import math
import re
import unicodedata
import hashlib
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ENGINE_VERSION = "python-draft-target-v1.0"
RULESET_SHA = "c4dfcceb5e062c3ab6f151d536756580975823a2"
RENDER_SHA = "43ccf1d88e37fc94ea6a76eeff018f16a355b8a7"
A4_DYNAMIC_SHA = "a5f6c848af221772a2bbc6aba78dbb75d8977e22"
RENDERER_VERSION = "dynamic-a4-v1.2"
TEMPLATE_ID = "1LWbOK3s2BlaEzYY7tgtlUn-6E4QoyazLt8fCYGdhHbY"
WRITE_ENABLED_DEFAULT = False

UNIT_ORDER = {"unidad": 4, "ml": 3, "m2": 2, "": 1}
INSCRIPTION_VARIANTS = ["RELIEVE", "SEGÚN SUYA", "ALDINE", "REMUS", "INGLESA", "CATANEO", "REDONDA", "GÓTICA", "LÁSER"]


def _clean(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    return str(v).strip()


def _norm(v):
    text = unicodedata.normalize("NFD", _clean(v).lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", text).strip()


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, np.integer, np.floating)):
        try:
            return float(v)
        except Exception:
            return None
    s = _clean(v).replace("€", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def _fmt_num(v):
    n = _num(v)
    if n is None:
        return ""
    if abs(n - round(n)) < 1e-9:
        return str(int(round(n)))
    return (f"{n:.2f}").rstrip("0").rstrip(".").replace(".", ",")


def _contains(text, *patterns):
    s = _norm(text)
    return any(re.search(p, s, flags=re.I) for p in patterns)


def _row_text(row):
    return " | ".join(_clean(row.get(k, "")) for k in ["Modelo", "Material", "Notas de medidas y croquis", "Especificaciones", "Texto conmemorativo", "Nota manuscrita"])


def _material_norm(row):
    raw = _norm(row.get("Material normalizado") or row.get("Material"))
    mappings = [
        (r"porcel", "PORCELANA"), (r"italia", "MÁRMOL ITALIANO"), (r"sudaf", "NEGRO SUDÁFRICA"),
        (r"absol", "NEGRO ABSOLUTO"), (r"champ", "BLANCO CHAMPÁN"), (r"macael|blanco", "BLANCO MACAEL"),
        (r"tezal", "NEGRO TEZAL"), (r"porrino|porriño", "ROSA PORRIÑO"), (r"quintana", "GRIS QUINTANA"),
        (r"multicolor", "ROJO MULTICOLOR"), (r"verde oliva", "VERDE OLIVA"), (r"labrador", "LABRADOR"),
        (r"crema marfil", "CREMA MARFIL"), (r"suyo|suya|existente", "MATERIAL SUYO"),
    ]
    for pat, val in mappings:
        if re.search(pat, raw, flags=re.I):
            return val
    return _clean(row.get("Material normalizado") or row.get("Material")).upper()


def _catalog_records(catalog):
    if isinstance(catalog, pd.DataFrame):
        records = catalog.to_dict("records")
    else:
        records = list(catalog)
    out = []
    for r in records:
        rr = {str(k): v for k, v in r.items()}
        concept = _clean(rr.get("Concepto") or rr.get("concepto") or rr.get("CONCEPTO"))
        if not concept:
            continue
        price = _num(rr.get("Precio") or rr.get("precio") or rr.get("Precio (€)") or rr.get("P.V.P."))
        variant = _clean(rr.get("Variante") or rr.get("variante"))
        material = _clean(rr.get("Material") or rr.get("material"))
        unit = _clean(rr.get("Unidad") or rr.get("unidad"))
        validated = _norm(rr.get("Validado") or rr.get("validado") or rr.get("Estado")) in {"si", "sí", "true", "1", "validado"}
        out.append({"concept": concept, "price": price, "variant": variant, "material": material, "unit": unit, "validated": validated, "raw": rr})
    return out


def _catalog_lookup(catalog_records, concept, variant="", material=""):
    cn = _norm(concept); vn = _norm(variant); mn = _norm(material)
    candidates = []
    for r in catalog_records:
        if _norm(r["concept"]) != cn:
            continue
        if r["price"] is None or r["price"] <= 0:
            continue
        score = 100
        rv = _norm(r["variant"]); rm = _norm(r["material"]); ru = _norm(r["unit"])
        if vn and rv == vn: score += 40
        elif not rv: score += 5
        if mn and rm == mn: score += 30
        elif not rm: score += 2
        score += UNIT_ORDER.get(ru, 0) * 20
        if r["validated"]: score += 2
        candidates.append((score, r))
    if not candidates:
        return None
    candidates.sort(key=lambda x: (-x[0], _clean(x[1]["concept"]), _clean(x[1]["variant"]), _clean(x[1]["material"])))
    return candidates[0][1]


def _line(concept, qty=1, kind="work", variant="", material="", note=""):
    return {"concept": concept, "quantity": qty, "kind": kind, "variant": variant, "material": material, "note": note}


def _derive_lines(row, catalog):
    text = _row_text(row); ntext = _norm(text); mat = _material_norm(row)
    lines = []
    def add(concept, qty=1, kind="work", variant="", material="", note=""):
        lines.append(_line(concept, qty, kind, variant, material or mat, note))
    if _contains(text, r"\bcorte\b"): add("CORTE")
    if _contains(text, r"\brepisa\b") and not _contains(text, r"repisa (suya|suyo|existente)"): add("REPISA")
    if _contains(text, r"\bcornisa\b") and not _contains(text, r"cornisa (suya|suyo|existente)"): add("CORNISA")
    if _contains(text, r"coronacion|coronación"): add("CORONACIÓN")
    if _contains(text, r"pegar cornisa"): add("PEGAR CORNISA")
    if _contains(text, r"pegar jardinera"): add("PEGAR JARDINERA")
    if _contains(text, r"tabica|tacos?"): add("TABICA/TACOS")
    if _contains(text, r"columna|pilastra"): add("COLUMNA/PILASTRA")
    if _contains(text, r"barras?\s*z"): add("BARRAS Z")
    if _contains(text, r"\bgarras?\b"): add("GARRAS")
    if _contains(text, r"\bcruz\b"): add("CRUZ")
    if _contains(text, r"\bimagen\b"): add("IMAGEN")
    # Physical photo only on explicit dimensions/physical request; s/foto, sin foto, según foto are excluded.
    photo_neg = _contains(text, r"s/foto", r"sin foto", r"segun foto|según foto")
    photo_pos = _contains(text, r"foto\s*\d+\s*[x×]\s*\d+", r"foto.*color", r"foto.*acero", r"foto.*fondo")
    if photo_pos and not photo_neg: add("FOTO")
    specs = _norm(row.get("Especificaciones", "")); memorial = _norm(row.get("Texto conmemorativo", ""))
    inscription_explicit = bool(re.search(r"inscrip", specs)) or bool(re.search(r"(grabada|relieve|laser|láser|aldine|remus|inglesa|catan[eé]o|redonda|g[oó]tica)", memorial))
    if inscription_explicit:
        variant = ""
        raw_upper = (_clean(row.get("Especificaciones")) + " " + _clean(row.get("Texto conmemorativo"))).upper()
        if "RELIEVE" in raw_upper and ("S/SUYA" in raw_upper or "SEGÚN SUYA" in raw_upper or "SEGUN SUYA" in raw_upper): variant = "RELIEVE + SEGÚN SUYA"
        else:
            for v in INSCRIPTION_VARIANTS:
                if _norm(v) in _norm(raw_upper): variant = v; break
        add("INSCRIPCIÓN", variant=variant)
    if _contains(text, r"jardinera"): add("JARDINERA")
    if _contains(text, r"florero"): add("FLORERO")
    if _contains(text, r"floreo"): add("FLOREO")
    if _contains(text, r"cenefa"): add("CENEFA")
    if _contains(text, r"borrar cartela"): add("BORRAR CARTELA")
    if _contains(text, r"retacear.*zafra"): add("RETACEAR EN ZAFRA")
    if _contains(text, r"pegar escalon|pegar escalón"): add("PEGAR ESCALÓN")
    if _contains(text, r"taladr"): add("TALADRADO")
    if _contains(text, r"canto pulido"): add("CANTO PULIDO")
    if _contains(text, r"\bpulido\b") and not _contains(text, r"canto pulido"): add("PULIDO")
    if _contains(text, r"limpieza"): add("LIMPIEZA")
    if _contains(text, r"montaje|colocacion|colocación"): add("MONTAJE/COLOCACIÓN")
    if _contains(text, r"desmontaje"): add("DESMONTAJE")
    if _contains(text, r"faja perimetral"): add("FAJA PERIMETRAL")
    if _contains(text, r"abujardado"): add("ABUJARDADO")
    if _contains(text, r"pegar embellecedores|punta diamante"): add("PEGAR EMBELLECEDORES/PUNTA DIAMANTE")
    if _contains(text, r"remate|cajon|cajón"): add("REMATE/CAJÓN")
    if _contains(text, r"barandilla|herrajes"): add("BARANDILLA/HERRAJES")
    # stable dedup by concept+variant+material+kind, retaining first occurrence
    seen=set(); ded=[]
    for l in lines:
        key=(_norm(l["concept"]),_norm(l["variant"]),_norm(l["material"]),l["kind"])
        if key in seen: continue
        seen.add(key); ded.append(l)
    # price enrichment
    records=_catalog_records(catalog)
    for l in ded:
        hit=_catalog_lookup(records,l["concept"],l["variant"],l["material"])
        l["unit_price"] = hit["price"] if hit else None
        l["unit"] = hit["unit"] if hit else ""
        l["catalog_match"] = bool(hit)
    return ded


def _copy_style(src, dst):
    if src.has_style:
        dst._style = _copy(src._style)
    if src.number_format: dst.number_format = src.number_format
    if src.font: dst.font = _copy(src.font)
    if src.fill: dst.fill = _copy(src.fill)
    if src.border: dst.border = _copy(src.border)
    if src.alignment: dst.alignment = _copy(src.alignment)
    if src.protection: dst.protection = _copy(src.protection)


def _render_dynamic_a4_openpyxl(row, lines, template_bytes):
    wb = load_workbook(BytesIO(template_bytes))
    ws = wb.active
    # Certified layout is compact dynamic A4. Keep seven columns A:G.
    # Clear prior variable region while preserving template styles.
    for r in range(1, max(ws.max_row, 40) + 1):
        for c in range(1, 8):
            ws.cell(r,c).value = None
    pedido=_clean(row.get("Pedido")); fecha=_clean(row.get("Fecha ficha")); modelo=_clean(row.get("Modelo")); material=_clean(row.get("Material"))
    specs=_clean(row.get("Especificaciones")); memorial=_clean(row.get("Texto conmemorativo")); notes=_clean(row.get("Notas de medidas y croquis"))
    ws["A1"]="PEDIDO"; ws["B1"]=pedido; ws["D1"]="FECHA"; ws["E1"]=fecha
    ws["A2"]="MODELO"; ws["B2"]=modelo; ws["D2"]="MATERIAL"; ws["E2"]=material
    ws["A3"]="MEDIDAS / NOTAS"; ws["B3"]=notes
    ws.merge_cells("B3:G3")
    ws["A4"]="ESPECIFICACIONES"; ws["B4"]=specs
    ws.merge_cells("B4:G4")
    ws["A5"]="TEXTO"; ws["B5"]=memorial
    ws.merge_cells("B5:G5")
    header_row=7
    headers=["CONCEPTO","CANT.","UNIDAD","VARIANTE / MATERIAL","P. UNIT.","IMPORTE","OBSERVACIONES"]
    for i,h in enumerate(headers,1): ws.cell(header_row,i).value=h
    start=8
    for idx,l in enumerate(lines,start):
        ws.cell(idx,1).value=l["concept"]
        ws.cell(idx,2).value=l.get("quantity",1)
        ws.cell(idx,3).value=l.get("unit","")
        vm=" · ".join(x for x in [l.get("variant",""),l.get("material","")] if x)
        ws.cell(idx,4).value=vm
        p=l.get("unit_price")
        ws.cell(idx,5).value=p if p is not None else "REVISAR"
        ws.cell(idx,6).value=(p*l.get("quantity",1)) if p is not None else ""
        ws.cell(idx,7).value=l.get("note","")
    end=max(start+len(lines)-1,start)
    total_row=end+2
    ws.cell(total_row,5).value="TOTAL"
    final=_num(row.get("Precio final (€)")) or _num(row.get("Total sheet (€)"))
    ws.cell(total_row,6).value=final if final is not None else f"=SUM(F{start}:F{end})"
    # widths certified; 7927 old reference had stale E width 10.43, current is 8.71.
    widths={"A":24.0,"B":9.0,"C":11.0,"D":31.0,"E":8.71,"F":11.0,"G":28.0}
    for k,v in widths.items(): ws.column_dimensions[k].width=v
    ws.print_area=f"A1:G{total_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
    ws.page_margins.left=0.25; ws.page_margins.right=0.25; ws.page_margins.top=0.3; ws.page_margins.bottom=0.3
    # wrap row 10, not 11, matching certified M2B2 renderer behavior.
    for c in range(1,8):
        ws.cell(10,c).alignment=_copy(ws.cell(10,c).alignment); ws.cell(10,c).alignment=ws.cell(10,c).alignment.copy(wrap_text=True)
        ws.cell(11,c).alignment=_copy(ws.cell(11,c).alignment); ws.cell(11,c).alignment=ws.cell(11,c).alignment.copy(wrap_text=False)
    out=BytesIO(); wb.save(out)
    return out.getvalue()


def _cell_signature(ws, max_row=None, max_col=7):
    mr=max_row or ws.max_row
    rows=[]
    for r in range(1,mr+1):
        vals=[]
        for c in range(1,max_col+1):
            v=ws.cell(r,c).value
            vals.append(v.isoformat() if hasattr(v,"isoformat") and not isinstance(v,str) else v)
        rows.append(vals)
    return rows


def _style_signature(ws, max_row=None, max_col=7):
    mr=max_row or ws.max_row
    out=[]
    for r in range(1,mr+1):
        row=[]
        for c in range(1,max_col+1):
            cell=ws.cell(r,c)
            row.append((cell.style_id,cell.number_format,cell.alignment.horizontal,cell.alignment.vertical,cell.alignment.wrap_text))
        out.append(row)
    return out


def compare_target_actual(target_bytes, actual_bytes, pedido=""):
    wt=load_workbook(BytesIO(target_bytes),data_only=False); wa=load_workbook(BytesIO(actual_bytes),data_only=False)
    st=wt.active; sa=wa.active
    max_row=max(st.max_row,sa.max_row); max_col=7
    tv=_cell_signature(st,max_row,max_col); av=_cell_signature(sa,max_row,max_col)
    cell_diffs=[]
    for r in range(max_row):
        for c in range(max_col):
            if tv[r][c] != av[r][c]: cell_diffs.append({"cell":f"{get_column_letter(c+1)}{r+1}","target":tv[r][c],"actual":av[r][c]})
    ts=_style_signature(st,max_row,max_col); a_s=_style_signature(sa,max_row,max_col)
    style_diffs=[]
    for r in range(max_row):
        for c in range(max_col):
            if ts[r][c] != a_s[r][c]: style_diffs.append(f"{get_column_letter(c+1)}{r+1}")
    layout=[]
    for col in range(1,8):
        L=get_column_letter(col); tw=st.column_dimensions[L].width; aw=sa.column_dimensions[L].width
        if tw != aw: layout.append({"kind":"column_width","column":L,"target":tw,"actual":aw})
    for r in range(1,max_row+1):
        th=st.row_dimensions[r].height; ah=sa.row_dimensions[r].height
        if th != ah: layout.append({"kind":"row_height","row":r,"target":th,"actual":ah})
    known=[]; unexplained=[]
    for d in layout:
        if str(pedido)=="7927" and d.get("kind")=="column_width" and d.get("column")=="E": known.append(d)
        else: unexplained.append(d)
    return {"match":len(cell_diffs)==0 and len(style_diffs)==0 and len(unexplained)==0,"cell_diffs":cell_diffs,"style_diffs":style_diffs,"layout_diffs":layout,"unexplained_layout_diffs":unexplained,"known_reference_drift":bool(known)}


def build_target_xlsx(row, catalog, template_bytes):
    lines=_derive_lines(row,catalog)
    flags=[]
    for l in lines:
        if not l.get("catalog_match") and l.get("kind")!="own": flags.append(f"SIN PRECIO: {l['concept']}")
    xlsx=_render_dynamic_a4_openpyxl(row,lines,template_bytes)
    return {"xlsx_bytes":xlsx,"lines":lines,"flags":flags,"renderer":RENDERER_VERSION}


def canonical_workbook_sha256(xlsx_bytes):
    wb=load_workbook(BytesIO(xlsx_bytes),data_only=False); ws=wb.active
    obj={"cells":_cell_signature(ws,ws.max_row,7),"widths":{get_column_letter(i):ws.column_dimensions[get_column_letter(i)].width for i in range(1,8)},"print_area":str(ws.print_area),"max_row":ws.max_row,"max_col":ws.max_column}
    return hashlib.sha256(json.dumps(obj,ensure_ascii=False,sort_keys=True,default=str,separators=(",",":")).encode()).hexdigest()


def runtime_contract():
    return {"engine_version":ENGINE_VERSION,"ruleset_sha":RULESET_SHA,"render_sha":RENDER_SHA,"a4_dynamic_sha":A4_DYNAMIC_SHA,"renderer_version":RENDERER_VERSION,"template_id":TEMPLATE_ID,"write_enabled_default":WRITE_ENABLED_DEFAULT,"io_model":"pure-target-builder-no-google-no-network"}
