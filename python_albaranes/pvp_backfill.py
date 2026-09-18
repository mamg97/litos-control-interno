from __future__ import annotations

import argparse
import io
import json
import math
import os
import re
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

import openpyxl
import xlrd

import parity as p
import total_parser as tp

MASTER_ID = p.MASTER_ID
SHEET_NAME = p.SHEET_NAME

H = {
    "id": "Pedido",
    "amount": "Importe trabajo / Debe (€)",
    "observation": "Observación de conciliación",
    "invoice": "Archivo factura / albarán (XLSX)",
    "total": "Total sheet (€)",
}

PVP_FACTOR = Decimal("1.262")
TARGET_NOTE = (
    "PVP histórico: Precio final/P.V.P. incluye IVA 21% + R.E. 5,2%. "
    "Se usa TOTAL explícito si el albarán incluye impuestos; si no, se calcula "
    "base × 1,262. Cuando el albarán no es validable o no concuerda internamente, "
    "la base procede del Importe trabajo / Debe del estadillo."
)
ORDER_RE = re.compile(r"(?<!\d)(\d{4})(?!\d)")


@dataclass
class WorkbookPricing:
    internal_order: str
    sums: list[Decimal]
    totals: list[Decimal]
    has_iva: bool
    has_re: bool
    iva_rate: Decimal | None
    re_rate: Decimal | None


def clean(v: Any) -> str:
    return str("" if v is None else v).strip()


def norm(v: Any) -> str:
    text = unicodedata.normalize("NFD", clean(v))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", text).strip().lower()


def dec(v: Any) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, Decimal):
        return v
    if isinstance(v, (int, float)):
        if not math.isfinite(float(v)):
            return None
        return Decimal(str(v))
    raw = re.sub(r"[^0-9,.-]", "", clean(v))
    if not raw or clean(v).startswith("="):
        return None
    comma, dot = raw.rfind(","), raw.rfind(".")
    if comma != -1 and dot != -1:
        decimal = "," if comma > dot else "."
        raw = raw.replace("." if decimal == "," else ",", "")
        raw = raw.replace(decimal, ".")
    elif comma != -1:
        raw = raw.replace(",", ".")
    try:
        return Decimal(raw)
    except Exception:
        return None


def money(v: Decimal | None) -> float | None:
    if v is None or v <= 0:
        return None
    return float(v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def bool_env(name: str) -> bool:
    return os.environ.get(name, "").strip().lower() == "true"


def assert_write_allowed(confirm: str) -> None:
    if not bool_env("LITOS_FREE_ONLY"):
        raise RuntimeError("LITOS_FREE_ONLY must be true")
    if not bool_env("LITOS_PVP_WRITE_ENABLED"):
        raise RuntimeError("LITOS_PVP_WRITE_ENABLED must be true")
    if confirm != "PVP_BACKFILL_V1":
        raise RuntimeError("Explicit confirmation token required")


def extract_order_from_row(row: list[Any]) -> str:
    labels = {"num", "numero", "nº", "n°", "pedido", "pedido nº", "pedido n°", "pedido no"}
    for i, v in enumerate(row):
        n = norm(v)
        if n in labels or "pedido nº" in n or "pedido n°" in n:
            for candidate in row[i + 1 : i + 7]:
                m = ORDER_RE.search(clean(candidate))
                if m:
                    return m.group(1)
                if isinstance(candidate, (int, float)) and 1000 <= float(candidate) <= 9999:
                    return str(int(candidate))
    return ""


def first_numeric_right_xlsx(formula_book, value_book, sheet_name: str, row: int, col: int) -> Decimal | None:
    memo: dict[tuple[str, str], Decimal | None] = {}
    ws = formula_book[sheet_name]
    for c in range(col + 1, ws.max_column + 1):
        coord = ws.cell(row=row, column=c).coordinate
        value = tp._eval_cell(formula_book, value_book, sheet_name, coord, memo, set())
        if value is not None:
            return value
    return None


def first_numeric_right_list(row: list[Any], col: int) -> Decimal | None:
    for candidate in row[col + 1 :]:
        d = dec(candidate)
        if d is not None:
            return d
    return None


def parse_xlsx(content: bytes) -> WorkbookPricing:
    fb = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    vb = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    try:
        order = ""
        sums: list[Decimal] = []
        totals: list[Decimal] = []
        has_iva = False
        has_re = False
        iva_rate = None
        re_rate = None

        for ws in fb.worksheets:
            rows = [[cell.value for cell in row] for row in ws.iter_rows()]
            if not order:
                for row in rows:
                    order = extract_order_from_row(row)
                    if order:
                        break

            for row_idx, row in enumerate(rows, start=1):
                for col_idx, value in enumerate(row, start=1):
                    n = norm(value)
                    if not n:
                        continue
                    if n == "suma":
                        d = first_numeric_right_xlsx(fb, vb, ws.title, row_idx, col_idx)
                        if d is not None and d > 0:
                            sums.append(d)
                    elif n == "total":
                        d = first_numeric_right_xlsx(fb, vb, ws.title, row_idx, col_idx)
                        if d is not None and d > 0:
                            totals.append(d)
                    elif n in {"iva", "i.v.a.", "i.v.a"}:
                        has_iva = True
                        d = first_numeric_right_xlsx(fb, vb, ws.title, row_idx, col_idx)
                        if d is not None:
                            iva_rate = d
                    elif n in {"r.e.", "r.e", "re", "recargo equivalencia", "recargo de equivalencia"}:
                        has_re = True
                        d = first_numeric_right_xlsx(fb, vb, ws.title, row_idx, col_idx)
                        if d is not None:
                            re_rate = d

        return WorkbookPricing(order, sums, totals, has_iva, has_re, iva_rate, re_rate)
    finally:
        fb.close()
        vb.close()


def parse_xls(content: bytes) -> WorkbookPricing:
    wb = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        order = ""
        sums: list[Decimal] = []
        totals: list[Decimal] = []
        has_iva = False
        has_re = False
        iva_rate = None
        re_rate = None
        for ws in wb.sheets():
            rows = [ws.row_values(i) for i in range(ws.nrows)]
            if not order:
                for row in rows:
                    order = extract_order_from_row(row)
                    if order:
                        break
            for row in rows:
                for col, value in enumerate(row):
                    n = norm(value)
                    if n == "suma":
                        d = first_numeric_right_list(row, col)
                        if d is not None and d > 0:
                            sums.append(d)
                    elif n == "total":
                        d = first_numeric_right_list(row, col)
                        if d is not None and d > 0:
                            totals.append(d)
                    elif n in {"iva", "i.v.a.", "i.v.a"}:
                        has_iva = True
                        iva_rate = first_numeric_right_list(row, col)
                    elif n in {"r.e.", "r.e", "re", "recargo equivalencia", "recargo de equivalencia"}:
                        has_re = True
                        re_rate = first_numeric_right_list(row, col)
        return WorkbookPricing(order, sums, totals, has_iva, has_re, iva_rate, re_rate)
    finally:
        wb.release_resources()


def parse_workbook(content: bytes, name: str) -> WorkbookPricing:
    lower = name.lower()
    if lower.endswith(".xls") and not lower.endswith(".xlsx"):
        return parse_xls(content)
    return parse_xlsx(content)


def choose_base(pricing: WorkbookPricing, master_base: Decimal | None) -> tuple[Decimal | None, str]:
    # Old workshop books sometimes have several SUMA rows and one aggregate TOTAL.
    # Prefer a no-tax TOTAL as net base. Otherwise choose the SUMA closest to the
    # accounting amount in the estadillo.
    if pricing.totals and not (pricing.has_iva or pricing.has_re):
        return pricing.totals[-1], "base_total_albaran"
    if pricing.sums:
        if master_base is None:
            return pricing.sums[-1], "base_suma_albaran"
        candidate = min(pricing.sums, key=lambda x: abs(x - master_base))
        tolerance = max(Decimal("1.00"), abs(master_base) * Decimal("0.01"))
        if abs(candidate - master_base) <= tolerance:
            return candidate, "base_suma_albaran"
    return master_base, "base_estadillo"


def pvp_from_sources(order_id: str, master_base: Decimal | None, pricing: WorkbookPricing | None) -> tuple[float | None, str]:
    if pricing is not None and pricing.internal_order and pricing.internal_order != order_id:
        if master_base is None:
            return None, f"albarán no concordante (interno {pricing.internal_order}); sin base de estadillo"
        return money(master_base * PVP_FACTOR), (
            f"calculado desde base de estadillo × 1,262; albarán no usado porque internamente corresponde a {pricing.internal_order}"
        )

    if pricing is not None and (pricing.has_iva or pricing.has_re) and pricing.totals:
        total = money(pricing.totals[-1])
        if total is not None:
            return total, "TOTAL explícito del albarán con impuestos/recargo"

    base, source = choose_base(pricing or WorkbookPricing("", [], [], False, False, None, None), master_base)
    if base is None or base <= 0:
        return None, "sin base económica utilizable"

    return money(base * PVP_FACTOR), (
        ("calculado desde base neta del albarán × 1,262" if source != "base_estadillo"
         else "calculado desde Importe trabajo / Debe del estadillo × 1,262")
    )


def append_note(current: str, note: str) -> str:
    if note in current:
        return current
    return f"{current} · {note}" if current else note


def read_master(sheets):
    rows, header_index, columns = p.read_sheet_values(sheets)
    for header in H.values():
        if header not in columns:
            raise RuntimeError(f"Missing master header: {header}")
    body_rows = max(0, len(rows) - header_index - 1)
    first_body_row = header_index + 2
    links = p.read_link_column(sheets, columns[H["invoice"]], first_body_row, body_rows)
    return rows, header_index, columns, links


def build_plan():
    drive, sheets = p.build_services()
    rows, header_index, columns, links = read_master(sheets)

    plans: list[dict[str, Any]] = []
    stats = {
        "candidate_rows": 0,
        "pvp_from_explicit_total": 0,
        "pvp_from_validated_albaran_base": 0,
        "pvp_from_estadillo_base": 0,
        "albaran_mismatch_fallbacks": 0,
        "no_pvp_source": 0,
        "files_parsed": 0,
        "parse_errors": 0,
    }
    cache: dict[str, tuple[str, WorkbookPricing | None, str | None]] = {}

    for body_index, row_index in enumerate(range(header_index + 1, len(rows))):
        row = rows[row_index]
        order_id = clean(row[columns[H["id"]]] if columns[H["id"]] < len(row) else "")
        if not re.fullmatch(r"\d{4}", order_id):
            continue

        current_total = row[columns[H["total"]]] if columns[H["total"]] < len(row) else None
        if dec(current_total) is not None:
            continue

        master_base = dec(row[columns[H["amount"]]] if columns[H["amount"]] < len(row) else None)
        link = links[body_index] if body_index < len(links) else p.LinkCell("", "")
        file_id = p._extract_drive_id(link.url)
        if master_base is None and not file_id:
            continue

        stats["candidate_rows"] += 1
        pricing = None
        parse_error = None
        file_name = ""

        if file_id:
            if file_id not in cache:
                try:
                    meta = drive.files().get(
                        fileId=file_id,
                        fields="id,name,mimeType",
                        supportsAllDrives=True,
                    ).execute()
                    file_name = clean(meta.get("name"))
                    content = p._download_file(drive, file_id)
                    pricing = parse_workbook(content, file_name)
                    cache[file_id] = (file_name, pricing, None)
                    stats["files_parsed"] += 1
                except Exception as exc:
                    cache[file_id] = ("", None, str(exc))
                    stats["parse_errors"] += 1
            file_name, pricing, parse_error = cache[file_id]

        pvp, reason = pvp_from_sources(order_id, master_base, pricing)
        if pvp is None:
            stats["no_pvp_source"] += 1
            continue

        if "TOTAL explícito" in reason:
            stats["pvp_from_explicit_total"] += 1
        elif "base neta del albarán" in reason:
            stats["pvp_from_validated_albaran_base"] += 1
        else:
            stats["pvp_from_estadillo_base"] += 1
        if "no usado porque internamente" in reason:
            stats["albaran_mismatch_fallbacks"] += 1

        current_obs = clean(row[columns[H["observation"]]] if columns[H["observation"]] < len(row) else "")
        detail = f"PVP histórico {pvp:.2f} €: {reason}."
        if parse_error:
            detail += " El albarán enlazado no pudo interpretarse; se utilizó la base contable."
        obs = append_note(current_obs, detail)

        plans.append({
            "row_index_zero": row_index,
            "order_id": order_id,
            "pvp": pvp,
            "observation": obs,
            "note": TARGET_NOTE + " " + detail,
        })

    stats["rows_backfillable"] = len(plans)
    stats["write_operations"] = 0
    return drive, sheets, columns, plans, stats


def sheet_id(sheets) -> int:
    response = sheets.spreadsheets().get(
        spreadsheetId=MASTER_ID,
        fields="sheets(properties(sheetId,title))",
    ).execute()
    for sheet in response.get("sheets", []) or []:
        props = sheet.get("properties", {}) or {}
        if props.get("title") == SHEET_NAME:
            return int(props["sheetId"])
    raise RuntimeError("Pedidos sheetId not found")


def cell_request(sid: int, row_zero: int, col_zero: int, cell: dict, fields: str) -> dict:
    return {
        "updateCells": {
            "range": {
                "sheetId": sid,
                "startRowIndex": row_zero,
                "endRowIndex": row_zero + 1,
                "startColumnIndex": col_zero,
                "endColumnIndex": col_zero + 1,
            },
            "rows": [{"values": [cell]}],
            "fields": fields,
        }
    }


def dry_run():
    _drive, _sheets, _columns, _plans, stats = build_plan()
    print("PVP_HISTORICAL_BACKFILL_DRY_RUN_OK")
    print(json.dumps(stats, ensure_ascii=False, sort_keys=True))
    return stats


def sync(confirm: str):
    assert_write_allowed(confirm)
    _drive, sheets, columns, plans, stats = build_plan()
    sid = sheet_id(sheets)

    requests: list[dict] = []
    for plan in plans:
        requests.append(
            cell_request(
                sid,
                plan["row_index_zero"],
                columns[H["total"]],
                {
                    "userEnteredValue": {"numberValue": float(plan["pvp"])},
                    "userEnteredFormat": {
                        "numberFormat": {
                            "type": "NUMBER",
                            "pattern": '#,##0.00 [$€-es-ES]',
                        }
                    },
                    "note": plan["note"],
                },
                "userEnteredValue,userEnteredFormat.numberFormat,note",
            )
        )
        requests.append(
            cell_request(
                sid,
                plan["row_index_zero"],
                columns[H["observation"]],
                {"userEnteredValue": {"stringValue": plan["observation"]}},
                "userEnteredValue",
            )
        )

    batch_size = 120
    batches = 0
    for offset in range(0, len(requests), batch_size):
        sheets.spreadsheets().batchUpdate(
            spreadsheetId=MASTER_ID,
            body={"requests": requests[offset : offset + batch_size]},
        ).execute()
        batches += 1

    result = {
        "mode": "PVP_HISTORICAL_BACKFILL_SYNC",
        "rows_written": len(plans),
        "cells_written": len(requests),
        "write_batches": batches,
        **stats,
    }
    print("PVP_HISTORICAL_BACKFILL_OK")
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--dry-run", action="store_true")
    group.add_argument("--sync", action="store_true")
    parser.add_argument("--confirm", default="")
    args = parser.parse_args()

    if args.dry_run:
        dry_run()
        return 0
    sync(args.confirm)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
