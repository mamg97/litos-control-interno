from __future__ import annotations

from io import BytesIO
import unittest

import openpyxl

from source_item_guard import (apply_explicit_absences, explicit_absent_labels,
                               has_explicit_absence_conflict, has_unrequested_rows,\n                               has_stale_absence_warnings)


def workbook_bytes():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for idx, label in enumerate(
        ["CORONACIÓN", "COLUMNA / PILASTRA", "CRUZ", "FLORERO",
         "CORNISA", "TABICA / TACOS", "REPISA", "IMAGEN / FOTO"], start=1
    ):
        sheet.cell(idx, 1).value = label
        sheet.cell(idx, 5).value = 1
        sheet.cell(idx, 6).value = 25
        sheet.cell(idx, 7).value = f"=E{idx}*F{idx}"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class SourceAbsenceTests(unittest.TestCase):
    def test_removes_only_explicit_absences_and_preserves_present_items(self):
        source = {
            "Especificaciones": (
                "CORONACION: -. COLUMNAS: -. PILASTRAS: -. CRUZ: -. "
                "FLORERO: -. CORNISA: -. TABICA Y TACOS: -. "
                "REPISA: fijada. IMAGEN: presente. FOTO: -."
            )
        }
        original = workbook_bytes()
        changed, removed = apply_explicit_absences(original, source)
        self.assertEqual(
            set(removed),
            {"CORONACIÓN", "COLUMNA / PILASTRA", "CRUZ", "FLORERO",
             "CORNISA", "TABICA / TACOS"},
        )
        self.assertTrue(has_unrequested_rows(original, removed))
        self.assertFalse(has_unrequested_rows(changed, removed))
        workbook = openpyxl.load_workbook(BytesIO(changed))
        try:
            sheet = workbook.active
            for row_idx in range(1, 7):
                self.assertTrue(all(sheet.cell(row_idx, col).value is None for col in range(1, 8)))
            self.assertEqual(sheet.cell(7, 1).value, "REPISA")
            self.assertEqual(sheet.cell(7, 5).value, 1)
            self.assertEqual(sheet.cell(8, 1).value, "IMAGEN / FOTO")
        finally:
            workbook.close()

    def test_explicit_absence_with_merged_detail_cells(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet["A4"] = "FLORERO"
        sheet["B4"] = "Synthetic optional item"
        sheet["E4"] = 1
        sheet["F4"] = 25
        sheet["G4"] = "=E4*F4"
        sheet.merge_cells("B4:D4")
        before = BytesIO()
        workbook.save(before)
        workbook.close()

        changed, removed = apply_explicit_absences(
            before.getvalue(), {"Especificaciones": "FLORERO: -."}
        )
        self.assertEqual(removed, ["FLORERO"])
        verified = openpyxl.load_workbook(BytesIO(changed))
        try:
            ws = verified.active
            self.assertIn("B4:D4", {str(x) for x in ws.merged_cells.ranges})
            self.assertTrue(all(ws.cell(4, col).value is None for col in range(1, 8)))
        finally:
            verified.close()

    def test_warning_only_absence_is_persisted_without_charge_row(self):
        """An absent item must not remain in the warning when its row is already empty."""
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet["A7"] = (
            "REVISAR DATOS · sin precio automático: CORTE, CORNISA, TABICA / TACOS"
            " · faltan medidas para: CORNISA"
        )
        sheet["A10"] = "CORTE"
        sheet["A11"] = "CORNISA"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        original = buffer.getvalue()
        source = {
            "Especificaciones": "CORTE: MATERIAL X; CORNISA: MATERIAL Y; TABICA Y TACOS: -."
        }

        self.assertTrue(has_stale_absence_warnings(original, ["TABICA / TACOS"]))
        changed, affected = apply_explicit_absences(original, source)
        self.assertEqual(affected, ["TABICA / TACOS"])
        self.assertFalse(has_stale_absence_warnings(changed, affected))
        self.assertFalse(has_unrequested_rows(changed, affected))
        verified = openpyxl.load_workbook(BytesIO(changed))
        try:
            self.assertEqual(
                verified.active["A7"].value,
                "REVISAR DATOS · sin precio automático: CORTE, CORNISA"
                " · faltan medidas para: CORNISA",
            )
            self.assertEqual(verified.active["A10"].value, "CORTE")
            self.assertEqual(verified.active["A11"].value, "CORNISA")
        finally:
            verified.close()

        again, again_affected = apply_explicit_absences(changed, source)
        self.assertEqual(again, changed)
        self.assertEqual(again_affected, [])

    def test_existing_stale_warning_requires_repair_even_when_target_is_clean(self):
        """The source, not target mutation labels, decides whether an old draft is stale."""
        source = {"Especificaciones": "CORTE: MATERIAL X; TABICA Y TACOS: -."}
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet["A7"] = "REVISAR DATOS · sin precio automático: CORTE, TABICA / TACOS"
        sheet["A10"] = "CORTE"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        existing = buffer.getvalue()

        clean_target, affected = apply_explicit_absences(existing, source)
        self.assertEqual(affected, ["TABICA / TACOS"])
        self.assertEqual(explicit_absent_labels(source), ["TABICA / TACOS"])
        self.assertTrue(has_explicit_absence_conflict(existing, source))
        self.assertFalse(has_explicit_absence_conflict(clean_target, source))
        self.assertFalse(has_explicit_absence_conflict(existing, {
            "Especificaciones": "TABICA Y TACOS: existente."
        }))
        self.assertFalse(has_explicit_absence_conflict(existing, {
            "Especificaciones": "TABICA Y TACOS: -; TABICA Y TACOS: existente."
        }))
        again, again_affected = apply_explicit_absences(clean_target, source)
        self.assertEqual(again, clean_target)
        self.assertEqual(again_affected, [])

    def test_missing_or_conflicting_absence_is_not_inferred(self):
        original = workbook_bytes()
        for specs in [
            "",
            "CRUZ: -; CRUZ: grabada.",
            "COLUMNAS: -. PILASTRAS: grabadas.",
            "FLORERO: 2 unidades.",
        ]:
            changed, removed = apply_explicit_absences(original, {"Especificaciones": specs})
            self.assertEqual(removed, [])
            self.assertEqual(changed, original)

    def test_absent_items_are_removed_from_warning_lists_and_regenerated(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet["A1"] = "CORNISA"
        sheet["E1"] = 1
        sheet["F1"] = 20
        sheet["A2"] = "TABICA / TACOS"
        sheet["E2"] = 1
        sheet["F2"] = 20
        sheet["A7"] = (
            "REVISAR DATOS · sin precio automático: CORNISA, TABICA / TACOS"
            " · faltan medidas para: REPISA, CORNISA, TABICA / TACOS"
        )
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        original = buffer.getvalue()
        absent = ["CORNISA", "TABICA / TACOS"]
        self.assertTrue(has_stale_absence_warnings(original, absent))
        changed, removed = apply_explicit_absences(
            original, {"Especificaciones": "CORNISA: -. TABICA Y TACOS: -."}
        )
        self.assertEqual(removed, absent)
        self.assertFalse(has_unrequested_rows(changed, absent))
        self.assertFalse(has_stale_absence_warnings(changed, absent))
        verified = openpyxl.load_workbook(BytesIO(changed))
        try:
            self.assertEqual(verified.active["A7"].value, "REVISAR DATOS · faltan medidas para: REPISA")
        finally:
            verified.close()
        again, again_removed = apply_explicit_absences(
            changed, {"Especificaciones": "CORNISA: -. TABICA Y TACOS: -."}
        )
        self.assertEqual(again, changed)
        self.assertEqual(again_removed, [])

    def test_explicit_quantity_is_not_erased(self):
        from runner import _apply_declared_item_quantities
        original = workbook_bytes()
        for specifications in ["FLORERO: 2 unidades.", "FLORERO: 2 floreros de referencia."]:
            changed, corrections = _apply_declared_item_quantities(
                original, {"Especificaciones": specifications}
            )
            self.assertEqual(len(corrections), 1)
            workbook = openpyxl.load_workbook(BytesIO(changed))
            try:
                self.assertEqual(workbook.active.cell(4, 5).value, 2)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
