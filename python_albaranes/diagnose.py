from __future__ import annotations

import io
import json
import re
from typing import Any

import openpyxl
import xlrd
from openpyxl.utils.cell import range_boundaries

import parity as p
from total_parser import SUM_RANGE_RE, read_total_from_bytes


def _safe_formula(value: Any) -> str:
    text = p._clean(value)
    if not text.startswith("="):
        return ""
    text = re.sub(r'"[^"]*"', '"…"', text)
    return text[:200]


def _sum_operands(formula: str, formula_sheet, value_book) -> list[dict]:
    compact = formula.replace(" ", "")
    match = SUM_RANGE_RE.match(compact)
    if not match:
        return []
    quoted_sheet, plain_sheet, start_ref, end_ref = match.groups()
    target_title = quoted_sheet or plain_sheet or formula_sheet.title
    if target_title not in value_book.sheetnames:
        return []
    value_sheet = value_book[target_title]
    start_ref = start_ref.replace("$", "")
    end_ref = end_ref.replace("$", "")
    min_col, min_row, max_col, max_row = range_boundaries(f"{start_ref}:{end_ref}")
    formula_target = formula_sheet.parent[target_title]
    out = []
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            fcell = formula_target.cell(row=row, column=col)
            vcell = value_sheet.cell(row=row, column=col)
            item = {
                "coord": fcell.coordinate,
                "cached_value": vcell.value,
                "number_format": fcell.number_format,
                "formula_type": fcell.data_type,
            }
            formula_text = _safe_formula(fcell.value)
            if formula_text:
                item["formula"] = formula_text
            elif fcell.value is not None:
                item["literal_value"] = fcell.value
            out.append(item)
    return out


def _inspect_openxml(content: bytes) -> dict:
    formula_book = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    value_book = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    try:
        hits = []
        for formula_sheet in formula_book.worksheets:
            value_sheet = value_book[formula_sheet.title]
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if p._normalize(cell.value) != "total":
                        continue
                    adjacent = []
                    for offset in range(1, 7):
                        col = cell.column + offset
                        formula_cell = formula_sheet.cell(row=cell.row, column=col)
                        value_cell = value_sheet.cell(row=cell.row, column=col)
                        if formula_cell.value is None and value_cell.value is None:
                            continue
                        item = {
                            "coord": formula_cell.coordinate,
                            "formula_type": formula_cell.data_type,
                            "cached_value": value_cell.value,
                            "number_format": formula_cell.number_format,
                        }
                        formula = _safe_formula(formula_cell.value)
                        if formula:
                            item["formula"] = formula
                            operands = _sum_operands(formula_cell.value, formula_sheet, value_book)
                            if operands:
                                item["sum_operands"] = operands
                        elif formula_cell.value is not None:
                            number = p._number(formula_cell.value)
                            item["literal_number"] = number
                            item["literal_type"] = type(formula_cell.value).__name__
                        adjacent.append(item)
                    hits.append(
                        {
                            "sheet": formula_sheet.title[:80],
                            "total_label_coord": cell.coordinate,
                            "adjacent": adjacent,
                        }
                    )
        calc = getattr(formula_book, "calculation", None)
        calc_info = {}
        if calc is not None:
            for attr in ("calcMode", "fullCalcOnLoad", "forceFullCalc", "calcOnSave"):
                value = getattr(calc, attr, None)
                if value is not None:
                    calc_info[attr] = value
        return {"format": "openxml", "total_hits": hits, "calculation": calc_info}
    finally:
        formula_book.close()
        value_book.close()


def _inspect_xls(content: bytes) -> dict:
    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        hits = []
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                values = sheet.row_values(row_index)
                for col, value in enumerate(values):
                    if p._normalize(value) != "total":
                        continue
                    adjacent = []
                    for offset, candidate in enumerate(values[col + 1 : col + 7], start=1):
                        if candidate in (None, ""):
                            continue
                        adjacent.append(
                            {
                                "col_offset": offset,
                                "value": candidate,
                                "number": p._number(candidate),
                                "type": type(candidate).__name__,
                            }
                        )
                    hits.append(
                        {
                            "sheet": sheet.name[:80],
                            "row_one_based": row_index + 1,
                            "col_one_based": col + 1,
                            "adjacent": adjacent,
                        }
                    )
        return {"format": "xls", "total_hits": hits}
    finally:
        book.release_resources()


def inspect_file(drive, file: p.DriveFile) -> dict:
    content = p._download_file(drive, file.file_id)
    lower = file.name.lower()
    if lower.endswith(".xls") and not lower.endswith(".xlsx"):
        return _inspect_xls(content)
    return _inspect_openxml(content)


def read_total_v2(drive, file: p.DriveFile) -> float | None:
    content = p._download_file(drive, file.file_id)
    return read_total_from_bytes(file.name, content)


def main() -> int:
    drive, sheets = p.build_services()
    index = p.scan_albaranes(drive)
    rows, header_index, columns = p.read_sheet_values(sheets)

    mismatches = []
    parsed_by_file: dict[str, float | None] = {}

    for row_index in range(header_index + 1, len(rows)):
        row = rows[row_index]
        id_col = columns[p.HEADERS["id"]]
        order_id = p._clean(row[id_col]) if id_col < len(row) else ""
        if not order_id or order_id not in index:
            continue

        entry = index[order_id]
        invoice = entry.get("invoice")
        draft = entry.get("draft")
        active = invoice or draft
        if active is None:
            continue

        if active.file_id not in parsed_by_file:
            parsed_by_file[active.file_id] = read_total_v2(drive, active)
        python_total = parsed_by_file[active.file_id]

        total_col = columns[p.HEADERS["total"]]
        sheet_raw = row[total_col] if total_col < len(row) else None
        sheet_total = p._number(sheet_raw)

        reason = ""
        if python_total is None:
            if sheet_total is not None or p._clean(sheet_raw):
                reason = "parser_unavailable_with_sheet_total"
        elif not p._same_number(sheet_raw, python_total):
            reason = "numeric_mismatch"

        if not reason:
            continue

        extension = active.name.lower().rsplit(".", 1)[-1] if "." in active.name else "unknown"
        detail = {
            "order_id": order_id,
            "active_kind": "invoice" if invoice is not None else "draft",
            "extension": extension,
            "reason": reason,
            "sheet_total": sheet_total,
            "python_total": python_total,
            "workbook": inspect_file(drive, active),
        }
        mismatches.append(detail)

    print("ALBARANES_DIAGNOSTIC_V2_OK")
    print(
        json.dumps(
            {
                "phase": "M6",
                "mode": "READ_ONLY_TOTAL_DIAGNOSTIC_V2",
                "mismatch_count": len(mismatches),
                "mismatches": mismatches,
                "write_operations": 0,
            },
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
