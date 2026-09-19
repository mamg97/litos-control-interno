from __future__ import annotations

import argparse
import io
import json
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Any

import openpyxl
import xlrd
from pypdf import PdfReader

import parity as p
from history_backfill import as_date, excel_serial, normalize

DELIVERY_HEADER = "Fecha entrega albarán"
ORDER_DATE_HEADER = "Fecha ficha"
OBSERVATION_HEADER = "Observación de conciliación"
INVOICE_HEADER = "Archivo factura / albarán (XLSX)"

SUPPORTED_EXTENSIONS = {".pdf", ".xls", ".xlsx", ".xlsm"}
ORDER_RE = re.compile(r"(?<!\d)(\d{4})(?!\d)")
PDF_ORDER_RE = re.compile(r"(?i)pedido\s*(?:n[º°o.]*)?\s*(\d{4})(?!\d)")
PDF_DELIVERY_RE = re.compile(
    r"(?im)^\s*fecha(?:\s+de\s+entrega)?\s*[.:]*\s*(?:\r?\n\s*)?"
    r"(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b"
)


@dataclass(frozen=True)
class CandidateFile:
    file_id: str
    name: str
    modified_time: str
    web_view_link: str


@dataclass(frozen=True)
class ParsedDocument:
    internal_order: str
    delivery_date: date | None
    raw_delivery_date: date | None
    parser: str


def clean(value: Any) -> str:
    return str(value if value is not None else "").strip()


def bool_env(name: str) -> bool:
    return os.environ.get(name, "").strip().lower() == "true"


def assert_write_allowed(confirm: str) -> None:
    if not bool_env("LITOS_FREE_ONLY"):
        raise RuntimeError("LITOS_FREE_ONLY must be true")
    if not bool_env("LITOS_DELIVERY_DATE_WRITE_ENABLED"):
        raise RuntimeError("LITOS_DELIVERY_DATE_WRITE_ENABLED must be true")
    if confirm != "DELIVERY_DATE_BACKFILL_V1":
        raise RuntimeError("Explicit confirmation token required")


def extension(name: str) -> str:
    lower = clean(name).lower()
    for ext in sorted(SUPPORTED_EXTENSIONS, key=len, reverse=True):
        if lower.endswith(ext):
            return ext
    return "." + lower.rsplit(".", 1)[-1] if "." in lower else ""


def is_draft(name: str) -> bool:
    return bool(p.DRAFT_RE.search(clean(name).lower()))


def list_children(drive, folder_id: str) -> list[dict]:
    return p._list_children(drive, folder_id)


def scan_definitive_documents(drive) -> tuple[dict[str, list[CandidateFile]], Counter[str]]:
    index: dict[str, list[CandidateFile]] = defaultdict(list)
    extensions: Counter[str] = Counter()
    visited: set[str] = set()
    stack = [p.ROOT_FOLDER_ID]

    while stack:
        folder_id = stack.pop()
        if folder_id in visited:
            continue
        visited.add(folder_id)
        for item in list_children(drive, folder_id):
            item_id = clean(item.get("id"))
            name = clean(item.get("name"))
            mime = clean(item.get("mimeType"))
            if mime == "application/vnd.google-apps.folder":
                if item_id == p.SYSTEM_FOLDER_ID or name == "_sistema":
                    continue
                stack.append(item_id)
                continue

            match = p.ORDER_ID_RE.search(name)
            if not match or is_draft(name):
                continue

            ext = extension(name)
            extensions[ext or "<sin_extension>"] += 1
            if ext not in SUPPORTED_EXTENSIONS:
                continue

            order_id = match.group(1)
            index[order_id].append(
                CandidateFile(
                    file_id=item_id,
                    name=name,
                    modified_time=clean(item.get("modifiedTime")),
                    web_view_link=clean(item.get("webViewLink")),
                )
            )

    for files in index.values():
        files.sort(key=lambda item: item.modified_time, reverse=True)
    return index, extensions


def read_master(sheets):
    response = sheets.spreadsheets().values().get(
        spreadsheetId=p.MASTER_ID,
        range=f"'{p.SHEET_NAME}'",
        valueRenderOption="UNFORMATTED_VALUE",
        dateTimeRenderOption="FORMATTED_STRING",
    ).execute()
    rows = response.get("values", [])
    header_index = next(
        (i for i, row in enumerate(rows) if any(clean(cell) == p.HEADERS["id"] for cell in row)),
        -1,
    )
    if header_index < 0:
        raise RuntimeError("Master header not found")
    headers = [clean(cell) for cell in rows[header_index]]
    columns = {header: i for i, header in enumerate(headers) if header}
    required = {
        p.HEADERS["id"],
        ORDER_DATE_HEADER,
        OBSERVATION_HEADER,
        INVOICE_HEADER,
        DELIVERY_HEADER,
    }
    missing = sorted(required - set(columns))
    if missing:
        raise RuntimeError("Master header contract changed: " + ", ".join(missing))
    return rows, header_index, columns


def link_column(sheets, column_index: int, first_body_row: int, body_rows: int):
    return p.read_link_column(sheets, column_index, first_body_row, body_rows)


def get_file_meta(drive, file_id: str) -> CandidateFile | None:
    try:
        item = drive.files().get(
            fileId=file_id,
            fields="id,name,modifiedTime,webViewLink,mimeType",
            supportsAllDrives=True,
        ).execute()
    except Exception:
        return None
    name = clean(item.get("name"))
    if not name or is_draft(name):
        return None
    return CandidateFile(
        file_id=clean(item.get("id")),
        name=name,
        modified_time=clean(item.get("modifiedTime")),
        web_view_link=clean(item.get("webViewLink")),
    )


def matrix_openxml(content: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            yield [list(row) for row in ws.iter_rows(values_only=True)], None
    finally:
        wb.close()


def matrix_xls(content: bytes):
    wb = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        for ws in wb.sheets():
            yield [ws.row_values(i) for i in range(ws.nrows)], wb.datemode
    finally:
        wb.release_resources()


def date_near_label(rows: list[list[Any]], row_index: int, col_index: int, datemode: int | None) -> date | None:
    positions: list[tuple[int, int]] = []
    for col in range(col_index + 1, min(len(rows[row_index]), col_index + 9)):
        positions.append((row_index, col))
    for r in range(row_index + 1, min(len(rows), row_index + 4)):
        for col in range(col_index, min(len(rows[r]), col_index + 9)):
            positions.append((r, col))
    for r, col in positions:
        candidate = as_date(rows[r][col], datemode=datemode)
        if candidate is not None:
            return candidate
    return None


def parse_workbook(content: bytes, name: str) -> ParsedDocument:
    lower = name.lower()
    matrices = matrix_xls(content) if lower.endswith(".xls") and not lower.endswith(".xlsx") else matrix_openxml(content)
    internal_order = ""
    date_candidates: list[date] = []

    for rows, datemode in matrices:
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                norm = normalize(value).rstrip(".:")
                if not norm:
                    continue
                if not internal_order and "pedido" in norm:
                    for candidate in row[col_index + 1 : col_index + 8]:
                        match = ORDER_RE.search(clean(candidate))
                        if match:
                            internal_order = match.group(1)
                            break
                        if isinstance(candidate, (int, float)) and not isinstance(candidate, bool):
                            number = int(candidate)
                            if 1000 <= number <= 9999:
                                internal_order = str(number)
                                break

                if norm in {"fecha", "fecha de entrega"}:
                    candidate = date_near_label(rows, row_index, col_index, datemode)
                    if candidate is not None:
                        date_candidates.append(candidate)

    raw = date_candidates[-1] if date_candidates else None
    return ParsedDocument(
        internal_order=internal_order,
        delivery_date=raw,
        raw_delivery_date=raw,
        parser="spreadsheet",
    )


def parse_pdf(content: bytes) -> ParsedDocument:
    reader = PdfReader(io.BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    order_match = PDF_ORDER_RE.search(text)
    dates = [as_date(value) for value in PDF_DELIVERY_RE.findall(text)]
    dates = [value for value in dates if value is not None]
    raw = dates[-1] if dates else None
    return ParsedDocument(
        internal_order=order_match.group(1) if order_match else "",
        delivery_date=raw,
        raw_delivery_date=raw,
        parser="pdf",
    )


def parse_document(drive, file: CandidateFile) -> ParsedDocument:
    content = p._download_file(drive, file.file_id)
    ext = extension(file.name)
    if ext == ".pdf":
        return parse_pdf(content)
    if ext in {".xls", ".xlsx", ".xlsm"}:
        return parse_workbook(content, file.name)
    raise RuntimeError("unsupported extension")


def plausible_delivery(candidate: date | None, order_date: date | None) -> bool:
    if candidate is None:
        return False
    today = datetime.now().date()
    if candidate < date(2000, 1, 1) or candidate > today + timedelta(days=1):
        return False
    if order_date is not None and candidate < order_date:
        return False
    return True


def append_observation(current: Any, note: str) -> str:
    base = clean(current)
    if not note or note in base:
        return base
    return note if not base else base + " · " + note


def build_plan(drive, sheets):
    rows, header_index, columns = read_master(sheets)
    body_rows = max(0, len(rows) - header_index - 1)
    first_body_row = header_index + 2
    invoice_links = link_column(
        sheets,
        columns[INVOICE_HEADER],
        first_body_row,
        body_rows,
    )
    scanned, ext_counts = scan_definitive_documents(drive)

    cache: dict[str, ParsedDocument | Exception] = {}
    plans: list[dict[str, Any]] = []
    stats: Counter[str] = Counter()
    conflicts: Counter[str] = Counter()

    for body_index, row_index in enumerate(range(header_index + 1, len(rows))):
        row = rows[row_index]
        id_col = columns[p.HEADERS["id"]]
        order_id = clean(row[id_col]) if id_col < len(row) else ""
        if not re.fullmatch(r"\d{4}", order_id):
            continue

        current = row[columns[DELIVERY_HEADER]] if columns[DELIVERY_HEADER] < len(row) else ""
        if clean(current):
            stats["already_filled"] += 1
            continue

        order_date_raw = row[columns[ORDER_DATE_HEADER]] if columns[ORDER_DATE_HEADER] < len(row) else None
        order_date = as_date(order_date_raw)

        candidates: list[CandidateFile] = []
        linked = invoice_links[body_index] if body_index < len(invoice_links) else p.LinkCell("", "")
        linked_id = p._extract_drive_id(linked.url)
        if linked_id:
            meta = get_file_meta(drive, linked_id)
            if meta is not None and extension(meta.name) in SUPPORTED_EXTENSIONS:
                candidates.append(meta)

        seen = {item.file_id for item in candidates}
        for item in scanned.get(order_id, []):
            if item.file_id not in seen:
                candidates.append(item)
                seen.add(item.file_id)

        if not candidates:
            stats["no_definitive_document"] += 1
            continue

        stats["rows_with_candidate"] += 1
        accepted: tuple[CandidateFile, ParsedDocument] | None = None
        row_conflicts: list[str] = []

        for candidate in candidates:
            stats[f"candidate_ext_{extension(candidate.name) or 'none'}"] += 1
            parsed_or_exc = cache.get(candidate.file_id)
            if parsed_or_exc is None:
                try:
                    parsed_or_exc = parse_document(drive, candidate)
                    cache[candidate.file_id] = parsed_or_exc
                    stats["documents_parsed"] += 1
                except Exception as exc:
                    parsed_or_exc = exc
                    cache[candidate.file_id] = exc
                    stats["parse_errors"] += 1

            if isinstance(parsed_or_exc, Exception):
                continue
            parsed = parsed_or_exc
            if parsed.internal_order and parsed.internal_order != order_id:
                conflicts["internal_id_mismatch"] += 1
                row_conflicts.append("ID interno no coincide")
                continue
            if parsed.raw_delivery_date is None:
                stats["documents_without_delivery_date"] += 1
                continue
            if not plausible_delivery(parsed.delivery_date, order_date):
                conflicts["implausible_delivery_date"] += 1
                row_conflicts.append(
                    "fecha de entrega documental incompatible "
                    + parsed.raw_delivery_date.isoformat()
                )
                continue
            accepted = (candidate, parsed)
            break

        if accepted is None:
            stats["rows_unresolved"] += 1
            if row_conflicts:
                current_obs = row[columns[OBSERVATION_HEADER]] if columns[OBSERVATION_HEADER] < len(row) else ""
                note = "Fecha entrega albarán no aplicada: " + "; ".join(sorted(set(row_conflicts)))
                updated_obs = append_observation(current_obs, note)
                if updated_obs != clean(current_obs):
                    plans.append({
                        "row_index_zero": row_index,
                        "order_id": order_id,
                        "delivery_date": None,
                        "observation": updated_obs,
                        "conflict_only": True,
                    })
            continue

        candidate, parsed = accepted
        current_obs = row[columns[OBSERVATION_HEADER]] if columns[OBSERVATION_HEADER] < len(row) else ""
        note = (
            "Fecha entrega albarán recuperada del documento definitivo "
            f"({extension(candidate.name).lstrip('.') or 'archivo'}; ID interno validado)"
        )
        plans.append({
            "row_index_zero": row_index,
            "order_id": order_id,
            "delivery_date": parsed.delivery_date,
            "observation": append_observation(current_obs, note),
            "conflict_only": False,
        })
        stats["dates_backfillable"] += 1

    summary = {
        "mode": "DELIVERY_DATE_BACKFILL_PLAN",
        "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
        "drive_candidate_extensions": dict(sorted(ext_counts.items())),
        "conflicts": dict(sorted(conflicts.items())),
        **dict(sorted(stats.items())),
        "planned_rows": len(plans),
        "write_operations": 0,
    }
    return rows, header_index, columns, plans, summary


def sheet_id(sheets) -> int:
    response = sheets.spreadsheets().get(
        spreadsheetId=p.MASTER_ID,
        fields="sheets(properties(sheetId,title))",
    ).execute()
    for item in response.get("sheets", []):
        props = item.get("properties", {})
        if props.get("title") == p.SHEET_NAME:
            return int(props["sheetId"])
    raise RuntimeError("Pedidos sheet not found")


def update_request(sheet_id_value: int, row_zero: int, col_zero: int, cell: dict, fields: str):
    return {
        "updateCells": {
            "range": {
                "sheetId": sheet_id_value,
                "startRowIndex": row_zero,
                "endRowIndex": row_zero + 1,
                "startColumnIndex": col_zero,
                "endColumnIndex": col_zero + 1,
            },
            "rows": [{"values": [cell]}],
            "fields": fields,
        }
    }


def sync(confirm: str):
    assert_write_allowed(confirm)
    drive, sheets = p.build_services()
    _rows, _header_index, columns, plans, before = build_plan(drive, sheets)
    sid = sheet_id(sheets)
    requests: list[dict] = []

    for plan in plans:
        if plan["delivery_date"] is not None:
            requests.append(
                update_request(
                    sid,
                    plan["row_index_zero"],
                    columns[DELIVERY_HEADER],
                    {
                        "userEnteredValue": {"numberValue": excel_serial(plan["delivery_date"])},
                        "userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "yyyy-mm-dd"}},
                    },
                    "userEnteredValue,userEnteredFormat.numberFormat",
                )
            )
        requests.append(
            update_request(
                sid,
                plan["row_index_zero"],
                columns[OBSERVATION_HEADER],
                {"userEnteredValue": {"stringValue": plan["observation"]}},
                "userEnteredValue",
            )
        )

    if requests:
        sheets.spreadsheets().batchUpdate(
            spreadsheetId=p.MASTER_ID,
            body={"requests": requests},
        ).execute()

    result = {
        "mode": "DELIVERY_DATE_BACKFILL_SYNC",
        "planned_rows_before_sync": before["planned_rows"],
        "dates_backfillable_before_sync": before.get("dates_backfillable", 0),
        "rows_written": len(plans),
        "write_operations": len(requests),
    }
    print("DELIVERY_DATE_BACKFILL_OK")
    print(json.dumps(result, sort_keys=True))
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true")
    mode.add_argument("--sync", action="store_true")
    parser.add_argument("--confirm", default="")
    args = parser.parse_args()

    drive, sheets = p.build_services()
    if args.dry_run:
        _rows, _header_index, _columns, _plans, summary = build_plan(drive, sheets)
        print("DELIVERY_DATE_BACKFILL_DRY_RUN_OK")
        print(json.dumps(summary, sort_keys=True))
        return 0

    sync(args.confirm)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
