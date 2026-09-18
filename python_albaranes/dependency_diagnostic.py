from __future__ import annotations

import io
import json
import re

import openpyxl

import parity as p
from total_parser import SUM_RANGE_RE, read_total_from_bytes

CELL_REF_RE = re.compile(r"(?<![A-Za-z0-9_])\$?([A-Z]{1,3})\$?(\d+)")


def _formula(value) -> str:
    text = p._clean(value)
    return text if text.startswith("=") else ""


def _sheet_and_range(formula: str, current_sheet: str):
    match = SUM_RANGE_RE.match(formula.replace(" ", ""))
    if not match:
        return None
    quoted, plain, start_ref, end_ref = match.groups()
    return quoted or plain or current_sheet, start_ref.replace("$", ""), end_ref.replace("$", "")


def _dependency_node(formula_book, value_book, sheet_name: str, coord: str, depth: int, seen: set[str]):
    key = f"{sheet_name}!{coord}"
    if depth > 8:
        return {"cell": key, "depth_limit": True}
    if key in seen:
        return {"cell": key, "cycle": True}

    seen = set(seen)
    seen.add(key)
    fsheet = formula_book[sheet_name]
    vsheet = value_book[sheet_name]
    fcell = fsheet[coord]
    vcell = vsheet[coord]
    node = {
        "cell": key,
        "formula_type": fcell.data_type,
        "formula": _formula(fcell.value),
        "literal_value": None if _formula(fcell.value) else fcell.value,
        "cached_value": vcell.value,
        "number_format": fcell.number_format,
    }

    formula = node["formula"]
    if not formula:
        return node

    dependencies = []
    sum_range = _sheet_and_range(formula, sheet_name)
    if sum_range:
        target_sheet, start_ref, end_ref = sum_range
        from openpyxl.utils.cell import range_boundaries

        min_col, min_row, max_col, max_row = range_boundaries(f"{start_ref}:{end_ref}")
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                dep_coord = fsheet.cell(row=row, column=col).coordinate if target_sheet == sheet_name else formula_book[target_sheet].cell(row=row, column=col).coordinate
                dependencies.append(_dependency_node(formula_book, value_book, target_sheet, dep_coord, depth + 1, seen))
    else:
        for col_letters, row_number in CELL_REF_RE.findall(formula):
            dep_coord = f"{col_letters}{row_number}"
            dependencies.append(_dependency_node(formula_book, value_book, sheet_name, dep_coord, depth + 1, seen))

    if dependencies:
        node["dependencies"] = dependencies
    return node


def _inspect_total_tree(content: bytes) -> dict:
    formula_book = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    value_book = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    try:
        trees = []
        for sheet in formula_book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if p._normalize(cell.value) != "total":
                        continue
                    for col in range(cell.column + 1, sheet.max_column + 1):
                        candidate = sheet.cell(row=cell.row, column=col)
                        if candidate.value is None:
                            continue
                        trees.append(
                            {
                                "total_label": f"{sheet.title}!{cell.coordinate}",
                                "value_cell": f"{sheet.title}!{candidate.coordinate}",
                                "tree": _dependency_node(
                                    formula_book,
                                    value_book,
                                    sheet.title,
                                    candidate.coordinate,
                                    0,
                                    set(),
                                ),
                            }
                        )
                        break
        return {"total_trees": trees}
    finally:
        formula_book.close()
        value_book.close()


def main() -> int:
    drive, sheets = p.build_services()
    index = p.scan_albaranes(drive)
    rows, header_index, columns = p.read_sheet_values(sheets)
    mismatches = []
    parsed = {}

    for row_index in range(header_index + 1, len(rows)):
        row = rows[row_index]
        id_col = columns[p.HEADERS["id"]]
        order_id = p._clean(row[id_col]) if id_col < len(row) else ""
        if not order_id or order_id not in index:
            continue
        entry = index[order_id]
        active = entry.get("invoice") or entry.get("draft")
        if active is None or not active.name.lower().endswith((".xlsx", ".xlsm")):
            continue
        if active.file_id not in parsed:
            content = p._download_file(drive, active.file_id)
            parsed[active.file_id] = (content, read_total_from_bytes(active.name, content))
        content, python_total = parsed[active.file_id]
        total_col = columns[p.HEADERS["total"]]
        sheet_raw = row[total_col] if total_col < len(row) else None
        sheet_total = p._number(sheet_raw)
        if python_total is not None and p._same_number(sheet_raw, python_total):
            continue
        if python_total is None and sheet_total is None and not p._clean(sheet_raw):
            continue
        mismatches.append(
            {
                "order_id": order_id,
                "sheet_total": sheet_total,
                "python_total": python_total,
                "dependency": _inspect_total_tree(content),
            }
        )

    print("ALBARANES_DEPENDENCY_DIAGNOSTIC_OK")
    print(json.dumps({"phase": "M6", "mismatch_count": len(mismatches), "mismatches": mismatches, "write_operations": 0}, ensure_ascii=False, sort_keys=True, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
