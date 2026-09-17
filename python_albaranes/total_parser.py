from __future__ import annotations

import io
import math
import re
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

import openpyxl
import xlrd
from openpyxl.utils.cell import range_boundaries

SUM_RANGE_RE = re.compile(
    r"^=SUM\((?:(?:'([^']+)'|([^'!]+))!)?(\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)\)$",
    re.IGNORECASE,
)
CELL_REF_RE = re.compile(r"^\$?([A-Z]+)\$?(\d+)$", re.IGNORECASE)
DIRECT_REF_RE = re.compile(r"^=\$?([A-Z]+)\$?(\d+)$", re.IGNORECASE)
BINARY_RE = re.compile(
    r"^=(\$?[A-Z]+\$?\d+|-?\d+(?:[.,]\d+)?)([+\-*/])(\$?[A-Z]+\$?\d+|-?\d+(?:[.,]\d+)?)$",
    re.IGNORECASE,
)
MAX_FORMULA_DEPTH = 24


def _clean(value: Any) -> str:
    return str(value if value is not None else "").strip()


def _normalize(value: Any) -> str:
    text = unicodedata.normalize("NFD", _clean(value))
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn").lower()


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None

    text = _clean(value)
    # Never interpret an Excel formula such as '=SUM(G11:G15)' as the digits
    # contained in its cell references.
    if text.startswith("="):
        return None

    raw = re.sub(r"[^0-9,.-]", "", text)
    if not raw:
        return None
    comma = raw.rfind(",")
    dot = raw.rfind(".")
    if comma != -1 and dot != -1:
        decimal = "," if comma > dot else "."
        raw = raw.replace("." if decimal == "," else ",", "")
        raw = raw.replace(decimal, ".")
    elif comma != -1:
        raw = raw.replace(",", ".")
    try:
        number = float(raw)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _decimal(value: Any) -> Decimal | None:
    if isinstance(value, Decimal):
        return value
    number = _number(value)
    if number is None:
        return None
    try:
        return Decimal(str(number))
    except InvalidOperation:
        return None


def _currency(value: Decimal | float | int | None) -> float | None:
    if value is None:
        return None
    decimal = value if isinstance(value, Decimal) else Decimal(str(value))
    if decimal <= 0:
        return None
    return float(decimal.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _cell_is_blank(formula_book, value_book, sheet_name: str, coord: str) -> bool:
    return (
        formula_book[sheet_name][coord].value in (None, "")
        and value_book[sheet_name][coord].value in (None, "")
    )


def _operand_value(
    token: str,
    formula_book,
    value_book,
    sheet_name: str,
    memo: dict[tuple[str, str], Decimal | None],
    stack: set[tuple[str, str]],
    depth: int,
) -> Decimal | None:
    token = token.strip()
    ref = CELL_REF_RE.match(token)
    if ref:
        coord = f"{ref.group(1).upper()}{ref.group(2)}"
        return _eval_cell(formula_book, value_book, sheet_name, coord, memo, stack, depth + 1)
    return _decimal(token)


def _eval_formula(
    formula: str,
    formula_book,
    value_book,
    sheet_name: str,
    memo: dict[tuple[str, str], Decimal | None],
    stack: set[tuple[str, str]],
    depth: int,
) -> Decimal | None:
    compact = formula.replace(" ", "")

    match = SUM_RANGE_RE.match(compact)
    if match:
        quoted_sheet, plain_sheet, start_ref, end_ref = match.groups()
        target_sheet = quoted_sheet or plain_sheet or sheet_name
        if target_sheet not in formula_book.sheetnames or target_sheet not in value_book.sheetnames:
            return None
        start_ref = start_ref.replace("$", "")
        end_ref = end_ref.replace("$", "")
        min_col, min_row, max_col, max_row = range_boundaries(f"{start_ref}:{end_ref}")
        total = Decimal("0")
        saw_value = False
        target = formula_book[target_sheet]
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                coord = target.cell(row=row, column=col).coordinate
                value = _eval_cell(
                    formula_book,
                    value_book,
                    target_sheet,
                    coord,
                    memo,
                    stack,
                    depth + 1,
                )
                if value is None:
                    if _cell_is_blank(formula_book, value_book, target_sheet, coord):
                        continue
                    return None
                total += value
                saw_value = True
        return total if saw_value else Decimal("0")

    match = BINARY_RE.match(compact)
    if match:
        left_token, operator, right_token = match.groups()
        left = _operand_value(left_token, formula_book, value_book, sheet_name, memo, stack, depth)
        right = _operand_value(right_token, formula_book, value_book, sheet_name, memo, stack, depth)
        if left is None or right is None:
            return None
        if operator == "+":
            return left + right
        if operator == "-":
            return left - right
        if operator == "*":
            return left * right
        if operator == "/":
            return None if right == 0 else left / right

    match = DIRECT_REF_RE.match(compact)
    if match:
        coord = f"{match.group(1).upper()}{match.group(2)}"
        return _eval_cell(formula_book, value_book, sheet_name, coord, memo, stack, depth + 1)

    return None


def _eval_cell(
    formula_book,
    value_book,
    sheet_name: str,
    coord: str,
    memo: dict[tuple[str, str], Decimal | None],
    stack: set[tuple[str, str]],
    depth: int = 0,
) -> Decimal | None:
    key = (sheet_name, coord)
    if key in memo:
        return memo[key]
    if depth > MAX_FORMULA_DEPTH or key in stack:
        return None

    formula_cell = formula_book[sheet_name][coord]
    cached_cell = value_book[sheet_name][coord]
    formula = formula_cell.value if formula_cell.data_type == "f" else None

    next_stack = set(stack)
    next_stack.add(key)

    result: Decimal | None = None
    if isinstance(formula, str):
        # Prefer deterministic evaluation of the small formula grammar actually
        # used by the workshop files. This avoids stale/missing cached formula
        # values and reproduces Google Sheets recalculation for these cases.
        result = _eval_formula(
            formula,
            formula_book,
            value_book,
            sheet_name,
            memo,
            next_stack,
            depth,
        )
        if result is None:
            result = _decimal(cached_cell.value)
    else:
        result = _decimal(formula_cell.value)
        if result is None:
            result = _decimal(cached_cell.value)

    memo[key] = result
    return result


def read_total_openxml(content: bytes) -> float | None:
    formula_book = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    value_book = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    try:
        memo: dict[tuple[str, str], Decimal | None] = {}
        for formula_sheet in formula_book.worksheets:
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if _normalize(cell.value) != "total":
                        continue

                    for column in range(cell.column + 1, formula_sheet.max_column + 1):
                        formula_cell = formula_sheet.cell(row=cell.row, column=column)
                        value = _eval_cell(
                            formula_book,
                            value_book,
                            formula_sheet.title,
                            formula_cell.coordinate,
                            memo,
                            set(),
                        )
                        if value is not None:
                            return _currency(value)
        return None
    finally:
        formula_book.close()
        value_book.close()


def read_total_xls(content: bytes) -> float | None:
    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        for sheet in workbook.sheets():
            for row_index in range(sheet.nrows):
                values = sheet.row_values(row_index)
                for col, value in enumerate(values):
                    if _normalize(value) != "total":
                        continue
                    for candidate in values[col + 1 :]:
                        decimal = _decimal(candidate)
                        if decimal is not None:
                            return _currency(decimal)
        return None
    finally:
        workbook.release_resources()


def read_total_from_bytes(file_name: str, content: bytes) -> float | None:
    lower = file_name.lower()
    if lower.endswith(".xls") and not lower.endswith(".xlsx"):
        return read_total_xls(content)
    return read_total_openxml(content)
