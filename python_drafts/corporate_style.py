from __future__ import annotations

from copy import copy
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

STYLE_VERSION = "corporate-a4-v2"
LOGO_SHA256 = "867ffb5417e24061d623e86e5f725bb3ec5b18fffd2fb81a6ff36ccd1829df84"
LOGO_PATH = Path(__file__).with_name("assets") / "company_logo.png"

NAVY = "16227A"
RED = "D9272E"
DARK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "EEF2F8"
LIGHT_GREY = "F5F6F7"
LINE = "D6DCE3"
AMBER = "FFF2CC"
AMBER_TEXT = "5C4A00"
PALE_RED = "FCE8E8"
RED_TEXT = "B42318"
WHITE = "FFFFFF"

THIN = Side(style="thin", color=LINE)
MEDIUM_NAVY = Side(style="medium", color=NAVY)
MEDIUM_RED = Side(style="medium", color=RED)
SOFT_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _text(value) -> str:
    return str(value if value is not None else "").strip()


def _find_row(ws, predicate, start: int = 1, end: int | None = None) -> int | None:
    end = end or ws.max_row
    for row in range(start, end + 1):
        if predicate(row):
            return row
    return None


def _cell_state(ws) -> tuple:
    out = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            out.append((cell.coordinate, cell.data_type, cell.value))
    return tuple(out)


def _apply_range_fill_font(ws, row: int, start_col: int, end_col: int, *, fill: str | None = None,
                           font: Font | None = None, alignment: Alignment | None = None,
                           border: Border | None = None) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if font:
            cell.font = copy(font)
        if alignment:
            cell.alignment = copy(alignment)
        if border:
            cell.border = copy(border)


def _verify_logo() -> None:
    if not LOGO_PATH.is_file():
        raise RuntimeError(f"Corporate logo asset missing: {LOGO_PATH}")
    actual = sha256(LOGO_PATH.read_bytes()).hexdigest()
    if actual != LOGO_SHA256:
        raise RuntimeError(f"Corporate logo SHA mismatch: expected {LOGO_SHA256}, got {actual}")


def apply_corporate_a4(xlsx_bytes: bytes) -> bytes:
    """Restyle a generated draft without changing any values or formulas.

    The semantic draft builder remains the source of truth. This function only
    changes presentation, embeds the exact approved logo, and configures A4
    print settings. It deliberately does not rewrite line descriptions,
    quantities, prices, formulas, warnings or observations.
    """
    _verify_logo()
    wb = load_workbook(BytesIO(xlsx_bytes), data_only=False)
    ws = wb.worksheets[0]
    before = _cell_state(ws)

    title_row = _find_row(ws, lambda r: _text(ws.cell(r, 1).value).upper() == "ALBARÁN")
    status_row = _find_row(ws, lambda r: _text(ws.cell(r, 1).value).upper().startswith("BORRADOR"))
    review_row = _find_row(ws, lambda r: _text(ws.cell(r, 1).value).upper().startswith("REVISAR DATOS"))
    order_row = _find_row(ws, lambda r: _text(ws.cell(r, 1).value).upper() == "PEDIDO Nº")
    material_row = _find_row(ws, lambda r: _text(ws.cell(r, 1).value).upper() == "MATERIAL", start=order_row or 1)
    table_header = _find_row(
        ws,
        lambda r: _text(ws.cell(r, 1).value).upper() == "CONCEPTO"
        and _text(ws.cell(r, 5).value).upper().startswith("CANT"),
        start=material_row or 1,
    )
    obs_row = _find_row(ws, lambda r: _text(ws.cell(r, 1).value).upper() == "OBS.", start=table_header or 1)
    footer_start = _find_row(
        ws,
        lambda r: _text(ws.cell(r, 1).value).upper() == "INSCRIPCIÓN"
        and _text(ws.cell(r, 5).value).upper() == "SUMA",
        start=(obs_row or 1) + 1,
    )
    total_row = _find_row(ws, lambda r: _text(ws.cell(r, 5).value).upper() == "TOTAL", start=footer_start or 1)

    required = {
        "title_row": title_row,
        "status_row": status_row,
        "order_row": order_row,
        "material_row": material_row,
        "table_header": table_header,
        "obs_row": obs_row,
        "footer_start": footer_start,
        "total_row": total_row,
    }
    missing = [name for name, value in required.items() if value is None]
    if missing:
        raise RuntimeError(f"Corporate style layout contract changed; missing {missing}")

    # Global document typography and neutral background.
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial", size=9, color=DARK, bold=cell.font.bold, italic=False)
            cell.alignment = copy(cell.alignment)
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="center",
                text_rotation=cell.alignment.text_rotation,
                wrap_text=cell.alignment.wrap_text,
                shrink_to_fit=cell.alignment.shrink_to_fit,
                indent=cell.alignment.indent,
            )
            cell.fill = PatternFill("solid", fgColor=WHITE)

    # Corporate header. Keep all existing business text exactly as generated.
    ws.row_dimensions[1].height = 44
    for row, height in ((2, 16), (3, 18), (4, 18), (5, 18), (title_row, 27)):
        ws.row_dimensions[row].height = height

    _apply_range_fill_font(
        ws, 1, 1, 7,
        font=Font(name="Arial", size=13, bold=True, italic=False, color=NAVY),
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    _apply_range_fill_font(
        ws, 2, 1, 7,
        font=Font(name="Arial", size=8.5, color=MUTED),
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    _apply_range_fill_font(ws, 3, 1, 7, font=Font(name="Arial", size=9, color=DARK))
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
    ws["F3"].alignment = Alignment(horizontal="right", vertical="center")
    _apply_range_fill_font(
        ws, 4, 1, 7,
        font=Font(name="Arial", size=9, bold=True, color=DARK),
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    _apply_range_fill_font(
        ws, 5, 1, 7,
        font=Font(name="Arial", size=9.5, bold=True, italic=False, color=NAVY),
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    _apply_range_fill_font(
        ws, title_row, 1, 7,
        font=Font(name="Arial", size=17, bold=True, color=NAVY),
        alignment=Alignment(horizontal="center", vertical="center"),
        border=Border(bottom=MEDIUM_RED),
    )

    logo = XLImage(str(LOGO_PATH))
    logo.width = 76
    logo.height = 49
    ws.add_image(logo, "A1")

    # Draft and review banners.
    ws.row_dimensions[status_row].height = 20
    _apply_range_fill_font(
        ws, status_row, 1, 7,
        fill=AMBER,
        font=Font(name="Arial", size=8.5, bold=True, color=AMBER_TEXT),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
    )
    if review_row:
        ws.row_dimensions[review_row].height = max(ws.row_dimensions[review_row].height or 20, 20)
        _apply_range_fill_font(
            ws, review_row, 1, 7,
            fill=PALE_RED,
            font=Font(name="Arial", size=8.5, bold=True, color=RED_TEXT),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        )

    # Order metadata block.
    for row in (order_row, material_row):
        ws.row_dimensions[row].height = 23
        _apply_range_fill_font(ws, row, 1, 7, fill=LIGHT_BLUE, border=SOFT_BORDER)
    for coordinate in (f"A{order_row}", f"C{order_row}", f"E{order_row}", f"A{material_row}", f"F{material_row}"):
        ws[coordinate].font = Font(name="Arial", size=8.5, bold=True, color=NAVY)
    ws.cell(order_row, 2).font = Font(name="Arial", size=10, bold=True, color=DARK)
    ws.cell(material_row, 2).font = Font(name="Arial", size=9.5, bold=True, color=DARK)

    # Main table.
    ws.row_dimensions[table_header].height = 23
    _apply_range_fill_font(
        ws, table_header, 1, 7,
        fill=NAVY,
        font=Font(name="Arial", size=8.5, bold=True, color=WHITE),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=Border(left=THIN, right=THIN, top=MEDIUM_NAVY, bottom=MEDIUM_NAVY),
    )

    for row in range(table_header + 1, obs_row):
        detail = _text(ws.cell(row, 2).value).upper()
        fill = WHITE
        concept_color = NAVY
        if "REVISAR PRECIO" in detail or "REVISAR MEDIDA" in detail:
            fill = AMBER
            concept_color = "7A4B00"
        elif "SIN CARGO AUTOMÁTICO" in detail:
            fill = LIGHT_GREY

        for col in range(1, 8):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = copy(SOFT_BORDER)
            if col == 1:
                cell.font = Font(name="Arial", size=8.8, bold=True, color=concept_color)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif 2 <= col <= 4:
                cell.font = Font(name="Arial", size=8.3, color=DARK)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.font = Font(name="Arial", size=8.7, color=DARK)
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # Observations.
    ws.row_dimensions[obs_row].height = max(ws.row_dimensions[obs_row].height or 28, 28)
    _apply_range_fill_font(
        ws, obs_row, 1, 1,
        fill=NAVY,
        font=Font(name="Arial", size=8.5, bold=True, color=WHITE),
        alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
        border=SOFT_BORDER,
    )
    _apply_range_fill_font(
        ws, obs_row, 2, 7,
        fill="FAFBFC",
        font=Font(name="Arial", size=8.1, color=DARK),
        alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
        border=SOFT_BORDER,
    )

    # Inscription and totals share the compact lower band.
    for row in range(footer_start, total_row + 1):
        for col in range(1, 8):
            ws.cell(row, col).border = copy(SOFT_BORDER)

    _apply_range_fill_font(
        ws, footer_start, 1, 1,
        fill=NAVY,
        font=Font(name="Arial", size=8.5, bold=True, color=WHITE),
        alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
    )
    # Style every merged cell in the inscription label block so Excel/Sheets renders consistently.
    for row in range(footer_start, total_row + 1):
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(row, 1).font = Font(name="Arial", size=8.5, bold=True, color=WHITE)
    for row in range(footer_start, total_row + 1):
        for col in range(2, 5):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=WHITE)
            ws.cell(row, col).font = Font(name="Arial", size=8.2, color=DARK)
            ws.cell(row, col).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for row in range(footer_start, total_row):
        for col in range(5, 7):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT_GREY)
        ws.cell(row, 5).font = Font(name="Arial", size=8.6, bold=True, color=NAVY)
        ws.cell(row, 6).font = Font(name="Arial", size=8.6, color=DARK)
        ws.cell(row, 7).font = Font(name="Arial", size=8.8, color=DARK)
        ws.cell(row, 7).alignment = Alignment(horizontal="right", vertical="center")

    for col in range(5, 8):
        ws.cell(total_row, col).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(total_row, col).font = Font(name="Arial", size=10, bold=True, color=WHITE)
        ws.cell(total_row, col).alignment = Alignment(horizontal="right" if col == 7 else "left", vertical="center")

    # Compact widths tuned for A4 portrait. Do not auto-fit: deterministic widths keep PDF output stable.
    widths = {"A": 17.5, "B": 19.0, "C": 17.5, "D": 17.5, "E": 10.5, "F": 13.5, "G": 13.5}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # One-page A4 print contract.
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.28
    ws.page_margins.bottom = 0.35
    ws.page_margins.header = 0.10
    ws.page_margins.footer = 0.18
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:G{ws.max_row}"
    ws.oddFooter.center.text = "BORRADOR · Revisar medidas, precios pendientes e incidencias antes de emitir."
    ws.oddFooter.center.font = "Arial"
    ws.oddFooter.center.size = 7
    ws.oddFooter.center.color = MUTED

    after = _cell_state(ws)
    if after != before:
        raise RuntimeError("Corporate style guard: cell values/formulas changed unexpectedly")

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _build_self_test_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.merge_cells("A1:G1")
    ws["A1"] = "Empresa de prueba"
    ws.merge_cells("A2:G2")
    ws["A2"] = "NIF"
    ws.merge_cells("A3:E3")
    ws["A3"] = "Dirección"
    ws.merge_cells("F3:G3")
    ws["F3"] = "Teléfono"
    ws.merge_cells("A4:G4")
    ws["A4"] = "LOCALIDAD"
    ws.merge_cells("A5:G5")
    ws["A5"] = "MÁRMOLES, PIEDRAS Y GRANITOS"
    ws.merge_cells("A6:G6")
    ws["A6"] = "ALBARÁN"
    ws.merge_cells("A7:G7")
    ws["A7"] = "BORRADOR · PRECIOS PROVISIONALES DEL CATÁLOGO OPERATIVO · REVISAR"
    ws.merge_cells("A8:G8")
    ws["A8"] = "REVISAR DATOS · sin precio automático: TEST"
    ws["A9"], ws["B9"], ws["C9"], ws["D9"], ws["E9"] = "PEDIDO Nº", 9999, "FECHA", "17/09/2026", "CONCEPTO"
    ws.merge_cells("F9:G9")
    ws["F9"] = "REFORMA"
    ws["A10"] = "MATERIAL"
    ws.merge_cells("B10:E10")
    ws["B10"] = "PIEDRA EXISTENTE"
    ws["F10"] = "PRECIO"
    ws["A11"], ws["B11"], ws["E11"], ws["F11"], ws["G11"] = "CONCEPTO", "DETALLE / MEDIDAS", "CANT.", "PRECIO UNIT.", "IMPORTE"
    ws.merge_cells("B11:D11")
    ws["A12"], ws["B12"], ws["E12"], ws["F12"] = "TEST", "PIEZA · [ud] · [REVISAR PRECIO]", 1, None
    ws.merge_cells("B12:D12")
    ws["A13"] = "OBS."
    ws.merge_cells("B13:G13")
    ws["B13"] = "Observación de prueba"
    ws["A14"] = "INSCRIPCIÓN"
    ws.merge_cells("A14:A17")
    ws.merge_cells("B14:D17")
    ws["B14"] = "D. E. P."
    ws["E14"] = "SUMA"
    ws["G14"] = "=SUM(G12:G12)"
    ws["E15"], ws["F15"], ws["G15"] = "IVA", 0.21, "=G14*F15"
    ws["E16"], ws["F16"], ws["G16"] = "R.E.", 0.052, "=G14*F16"
    ws["E17"], ws["G17"] = "TOTAL", "=SUM(G14:G16)"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def self_test() -> dict:
    styled = apply_corporate_a4(_build_self_test_workbook())
    wb = load_workbook(BytesIO(styled), data_only=False)
    ws = wb.active
    assert ws.page_setup.paperSize == ws.PAPERSIZE_A4
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 1
    assert ws.print_area == "'Hoja1'!$A$1:$G$17"
    assert len(ws._images) == 1
    assert ws["G14"].value == "=SUM(G12:G12)"
    assert ws["A11"].fill.fgColor.rgb.endswith(NAVY)
    return {
        "status": "CORPORATE_A4_SELF_TEST_OK",
        "style_version": STYLE_VERSION,
        "logo_sha256": LOGO_SHA256,
        "print_area": ws.print_area,
        "images": len(ws._images),
    }


if __name__ == "__main__":
    print(self_test())
